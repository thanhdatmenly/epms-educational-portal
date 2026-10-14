from django.shortcuts import render, redirect
from django.contrib.auth.decorators import login_required
from django.contrib import messages
from .models import TaiKhoan, GiaoVien,LopHoc,HocSinh,PhuHuynh,HocKy,PhuHuynhHocSinh,MonHoc,ToHop,BangDiem # Nhớ import các bảng cần dùng
import pandas as pd
from django.http import JsonResponse  # Thêm dòng này lên đầu file
from django.contrib.auth import authenticate, login,logout
from django.shortcuts import render, redirect
from django.views.decorators.csrf import ensure_csrf_cookie
from .models import XinNghiPhep, PhuHuynhHocSinh,DiemDanh
import json
from datetime import datetime,timedelta
import re
from django.http import JsonResponse
from .models import ThongBao
from django.views.decorators.http import require_POST
from django.http import JsonResponse
from .models import BangDiem, HocSinh
from .models import HoiThoai, TinNhan # Chắc chắn đã import
from django.db.models import Q
import io
from django.http import HttpResponse
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
# Import Model và Tool của chúng ta
from .models import HoiThoai, TinNhan, PhuHuynhHocSinh
from .tools import CongCuPhuHuynh
from google.genai import types
from .rag_engine import tra_cuu_noi_quy_truong
from google import genai
import os
# ================= HÀM XỬ LÝ ĐĂNG NHẬP CHÍNH =================
@ensure_csrf_cookie
def login_view(request):
    if request.method == 'POST':
        ten_dang_nhap = request.POST.get('username')
        mat_khau = request.POST.get('password')
        loai_form = request.POST.get('loai_form')  # 'hs', 'ph', hoặc 'gv'
        
        user = authenticate(request, username=ten_dang_nhap, password=mat_khau)
        
        if user is not None:
            di_dung_cua = False
            
            # --- CHỐT CHẶN MỚI: Định danh Admin tuyệt đối ---
            # Nếu là superuser thì chắc chắn là admin, mặc kệ cái 'vai_tro' bị gán default là gì
            la_admin = user.is_superuser or user.vai_tro == 'admin'
            
            # TRƯỜNG HỢP 1: Nếu là Admin -> BẮT BUỘC phải qua cửa Giáo viên ('gv')
            if la_admin:
                if loai_form == 'gv':
                    di_dung_cua = True
            
            # TRƯỜNG HỢP 2: Người dùng bình thường -> Phải đi đúng cửa của mình
            else:
                if user.vai_tro == loai_form:
                    di_dung_cua = True
                    
            # XỬ LÝ KẾT QUẢ RÀO CHẮN
            if not di_dung_cua:
                messages.error(request, "❌ Tài khoản không được phép đăng nhập ở cổng này!")
                return redirect('login')
            
            # CHO PHÉP ĐĂNG NHẬP VÀ ĐIỀU HƯỚNG
            login(request, user) 
            
            if la_admin:
                return redirect('admin_dashboard')  
            elif user.vai_tro == 'gv':
                return redirect('teacher_dashboard') 
            elif user.vai_tro == 'hs':
                return redirect('student_dashboard') 
            elif user.vai_tro == 'ph':
                return redirect('parent_dashboard')  
                
        else:
            messages.error(request, "❌ Tài khoản hoặc mật khẩu không chính xác!")
            return redirect('login')
            
    return render(request, 'login.html')

@login_required(login_url='login')
def admin_dashboard(request):
    # 1. Lấy toàn bộ danh sách giáo viên từ Database
    danh_sach_gv = GiaoVien.objects.all()
    danh_sach_lop = LopHoc.objects.all()
    danh_sach_hs = HocSinh.objects.all()
    danh_sach_ph = PhuHuynh.objects.all()
    danh_sach_monhoc = MonHoc.objects.all()
    danh_sach_hocky = HocKy.objects.all().order_by('-id')
    danh_sach_tohop = ToHop.objects.all()
    
    # =======================================================
    # THÊM ĐOẠN NÀY: ĐẾM SỐ ĐƠN XIN PHÉP ĐANG CHỜ DUYỆT TRÊN TOÀN TRƯỜNG
    # =======================================================
    from .models import XinNghiPhep
    so_don_cho_duyet = XinNghiPhep.objects.filter(trang_thai='Cho_Duyet').count()
    # Lệnh quét Database đếm xem có bao nhiêu học sinh
    tong_hoc_sinh = HocSinh.objects.count()
    # Lệnh đếm giáo viên (THÊM DÒNG NÀY VÀO)
    tong_giao_vien = GiaoVien.objects.count()

    # ===============================================
    # [THÊM MỚI] TÍNH CHUYÊN CẦN HÔM NAY VÀ SO SÁNH HÔM QUA
    # ===============================================
    hom_nay_date = datetime.now().date()
    hom_qua_date = hom_nay_date - timedelta(days=1)
    
    # 1. Tính Tỷ lệ hôm nay
    tong_dd_nay = DiemDanh.objects.filter(ngay=hom_nay_date).count()
    vang_nay = DiemDanh.objects.filter(ngay=hom_nay_date, trang_thai__in=['Vang_CP', 'Vang_KP']).count()
    ty_le_nay = round(((tong_dd_nay - vang_nay) / tong_dd_nay) * 100, 1) if tong_dd_nay > 0 else 100.0
    
    # 2. Tính Tỷ lệ hôm qua
    tong_dd_qua = DiemDanh.objects.filter(ngay=hom_qua_date).count()
    vang_qua = DiemDanh.objects.filter(ngay=hom_qua_date, trang_thai__in=['Vang_CP', 'Vang_KP']).count()
    ty_le_qua = round(((tong_dd_qua - vang_qua) / tong_dd_qua) * 100, 1) if tong_dd_qua > 0 else 100.0
    
    # 3. Tính chênh lệch
    chenh_lech = round(ty_le_nay - ty_le_qua, 1)
    chenh_lech_abs = abs(chenh_lech) # Lấy số dương để giao diện dễ hiển thị
    # ===============================================
    # [THÊM MỚI] ĐẾM SỐ VẮNG CỦA TOÀN TRƯỜNG
    # ===============================================
    # Lưu ý: Không lọc theo lớp nữa nên nó sẽ đếm toàn bộ Database
    tong_vang_cp_toan_truong = DiemDanh.objects.filter(trang_thai='Vang_CP').count()
    tong_vang_kp_toan_truong = DiemDanh.objects.filter(trang_thai='Vang_KP').count()
    # (Đoạn đếm vắng có phép/không phép toàn trường ban nãy của bác ở đây...)

    # ===============================================
    # [THÊM MỚI] TÍNH TỶ LỆ ĐI HỌC THEO TUẦN (LINE CHART TOÀN TRƯỜNG)
    # ===============================================
    ty_le_tuan_toan_truong = []
    tong_ty_le_tt = 0
    so_ngay_co_diem_danh_tt = 0
    thoi_gian_bieu_do_tt = "Chưa có dữ liệu"

    # Tìm ngày điểm danh mới nhất của TOÀN TRƯỜNG (Bỏ lọc lớp)
    ngay_moi_nhat_tt = DiemDanh.objects.order_by('-ngay').values_list('ngay', flat=True).first()

    if ngay_moi_nhat_tt:
        tuan_so_tt = ngay_moi_nhat_tt.isocalendar()[1]
        ngay_str_tt = ngay_moi_nhat_tt.strftime('%d/%m/%Y')
        thoi_gian_bieu_do_tt = f"Tuần {tuan_so_tt} — {ngay_str_tt}"

        thu_hai_tt = ngay_moi_nhat_tt - timedelta(days=ngay_moi_nhat_tt.weekday())

        for i in range(6):
            ngay_xet = thu_hai_tt + timedelta(days=i)
            # Đếm xem ngày đó toàn trường có bao nhiêu lượt điểm danh
            tong_luot_dd = DiemDanh.objects.filter(ngay=ngay_xet).count()

            if tong_luot_dd == 0:
                ty_le_tuan_toan_truong.append('null')
            else:
                # Đếm số học sinh vắng toàn trường ngày đó
                so_vang = DiemDanh.objects.filter(
                    ngay=ngay_xet,
                    trang_thai__in=['Vang_CP', 'Vang_KP']
                ).count()

                ty_le = round(((tong_luot_dd - so_vang) / tong_luot_dd) * 100, 1)
                ty_le_tuan_toan_truong.append(ty_le)
                tong_ty_le_tt += ty_le
                so_ngay_co_diem_danh_tt += 1
    else:
        hom_nay_tt = datetime.now()
        thoi_gian_bieu_do_tt = f"Tuần {hom_nay_tt.isocalendar()[1]} — {hom_nay_tt.strftime('%d/%m/%Y')}"

    chuoi_ty_le_tt = "[" + ",".join(str(x) for x in ty_le_tuan_toan_truong) + "]" if ty_le_tuan_toan_truong else "[100,100,100,100,100,100]"
    trung_binh_tuan_tt = round(tong_ty_le_tt / so_ngay_co_diem_danh_tt, 1) if so_ngay_co_diem_danh_tt > 0 else 100.0


    context = {
        'danh_sach_gv': danh_sach_gv,
        'danh_sach_lop': danh_sach_lop,
        'danh_sach_hs': danh_sach_hs,
        'danh_sach_ph': danh_sach_ph,
        'danh_sach_hocky': danh_sach_hocky,
        'danh_sach_monhoc': danh_sach_monhoc,
        'danh_sach_tohop': danh_sach_tohop,
        
        # BƠM BIẾN NÀY RA NGOÀI HTML
        'so_don_cho_duyet': so_don_cho_duyet, 
        'tong_hoc_sinh': tong_hoc_sinh,
        'tong_giao_vien': tong_giao_vien,
        # --- BÁC THÊM 3 DÒNG NÀY VÀO ĐÂY ---
        'ty_le_nay': ty_le_nay,
        'chenh_lech': chenh_lech,
        'chenh_lech_abs': chenh_lech_abs,
        # TRUYỀN 2 CON SỐ VỪA ĐẾM RA NGOÀI HTML
        'tong_vang_cp_toan_truong': tong_vang_cp_toan_truong,
        'tong_vang_kp_toan_truong': tong_vang_kp_toan_truong,
        'chuoi_ty_le_tt': chuoi_ty_le_tt,
        'trung_binh_tuan_tt': trung_binh_tuan_tt,
        'thoi_gian_bieu_do_tt': thoi_gian_bieu_do_tt,
    }
    return render(request, 'admin_dashboard.html', context)


from django.contrib.auth.decorators import login_required

# Bắt buộc phải đăng nhập mới được vào
@login_required(login_url='login')
def teacher_dashboard(request):
    # 1. Bảo vệ lớp 2: Chắc chắn người này là Giáo viên
    if request.user.vai_tro != 'gv':
        messages.error(request, "Bạn không có quyền truy cập trang này!")
        return redirect('login')

    # 2. Lấy hồ sơ Giáo viên đang đăng nhập (nhờ liên kết OneToOneField)
    try:
        gv_hien_tai = request.user.giao_vien
    except Exception:
        messages.error(request, "Tài khoản của bạn chưa được liên kết hồ sơ Giáo viên!")
        return redirect('login')

    # 3. Lấy lớp chủ nhiệm (nếu có)
    lop_chu_nhiem = gv_hien_tai.lop_chu_nhiem.all()
    lop_chinh = lop_chu_nhiem.first()  # Lớp chính (đầu tiên)

    # 4. Lấy danh sách học sinh của lớp chủ nhiệm
    danh_sach_hs = []
    if lop_chinh:
        danh_sach_hs = lop_chinh.danh_sach_hs.all().order_by('ho_ten')

    danh_sach_don = []
    so_don_cho_duyet = 0
    tong_vang_cp = 0
    tong_vang_kp = 0
    danh_sach_ph_lop = []
    buoi_da_diem_danh = []
    if lop_chinh:
        # Đếm số buổi vắng
        from .models import DiemDanh
        tong_vang_cp = DiemDanh.objects.filter(hoc_sinh__lop_hoc=lop_chinh, trang_thai='Vang_CP').count()
        tong_vang_kp = DiemDanh.objects.filter(hoc_sinh__lop_hoc=lop_chinh, trang_thai='Vang_KP').count()
        # ← THÊM MỚI: kiểm tra hôm nay CSDL đã có buổi nào điểm danh chưa
        from datetime import date
        buoi_da_diem_danh = list(
            DiemDanh.objects.filter(hoc_sinh__lop_hoc=lop_chinh, ngay=date.today())
            .values_list('buoi_hoc', flat=True).distinct()
        )
        # ===============================================
        # [THÊM MỚI] 3. TÍNH TỶ LỆ ĐI HỌC THEO TUẦN (LINE CHART)
        # ===============================================
        ty_le_tuan_thuc_te = []
        tong_ty_le = 0
        so_ngay_co_diem_danh = 0
        
        tong_hs = len(danh_sach_hs)
        if tong_hs > 0:
            ngay_moi_nhat = DiemDanh.objects.filter(hoc_sinh__lop_hoc=lop_chinh).order_by('-ngay').values_list('ngay', flat=True).first()
            if ngay_moi_nhat:
                # 2. PYTHON TỰ ĐỘNG DỊCH NGÀY THÀNH "Tuần thứ mấy trong năm"
                tuan_so = ngay_moi_nhat.isocalendar()[1]
                ngay_str = ngay_moi_nhat.strftime('%d/%m/%Y')
                thoi_gian_bieu_do = f"Tuần {tuan_so} — {ngay_str}"
                thu_hai = ngay_moi_nhat - timedelta(days=ngay_moi_nhat.weekday())
                
                for i in range(6): 
                    ngay_xet = thu_hai + timedelta(days=i)
                    
                    # Quét xem ngày đó CÓ BẤT KỲ AI ĐƯỢC ĐIỂM DANH KHÔNG
                    tong_luot_dd = DiemDanh.objects.filter(hoc_sinh__lop_hoc=lop_chinh, ngay=ngay_xet).count()
                    
                    if tong_luot_dd == 0:
                        # BÍ QUYẾT: Trả về chữ 'null' để JS hiểu là ngày tương lai, ngừng vẽ biểu đồ!
                        ty_le_tuan_thuc_te.append('null')
                    else:
                        so_vang = DiemDanh.objects.filter(
                            hoc_sinh__lop_hoc=lop_chinh,
                            ngay=ngay_xet,
                            trang_thai__in=['Vang_CP', 'Vang_KP']
                        ).count()
                        
                        # Tính % dựa trên số lượt điểm danh thực tế (Để lỡ 1 ngày điểm danh 2 buổi Sáng/Chiều không bị lố)
                        ty_le = round(((tong_luot_dd - so_vang) / tong_luot_dd) * 100, 1)
                        ty_le_tuan_thuc_te.append(ty_le)
                        tong_ty_le += ty_le
                        so_ngay_co_diem_danh += 1
            else:
                # Dự phòng: Nếu chưa điểm danh bao giờ thì lấy luôn ngày hôm nay
                from datetime import datetime
                hom_nay = datetime.now()
                tuan_so = hom_nay.isocalendar()[1]
                ngay_str = hom_nay.strftime('%d/%m/%Y')
                thoi_gian_bieu_do = f"Tuần {tuan_so} — {ngay_str}"
        # Đóng gói thành chuỗi JSON thô để HTML đọc mượt mà
        chuoi_ty_le = "[" + ",".join(str(x) for x in ty_le_tuan_thuc_te) + "]" if ty_le_tuan_thuc_te else "[100,100,100,100,100,100]"
        trung_binh_tuan = round(tong_ty_le / so_ngay_co_diem_danh, 1) if so_ngay_co_diem_danh > 0 else 100.0
        # Lấy tất cả đơn của học sinh trong lớp
        danh_sach_don = XinNghiPhep.objects.filter(
            hoc_sinh__lop_hoc=lop_chinh
        ).order_by('-ngay_tao').select_related('hoc_sinh', 'phu_huynh')

        so_don_cho_duyet = danh_sach_don.filter(trang_thai='Cho_Duyet').count()

        # Lấy danh sách phụ huynh của lớp (cho hộp thư)
        danh_sach_ph_lop = PhuHuynhHocSinh.objects.filter(
            hoc_sinh__lop_hoc=lop_chinh
        ).select_related('phu_huynh', 'hoc_sinh').order_by('hoc_sinh__ho_ten')
        # --- ĐOẠN MỚI THÊM: ĐẾM TIN NHẮN CHƯA ĐỌC ---
        danh_sach_ph_kem_tin_nhan = []
        tong_tin_nhan_chua_doc = 0

        for phhs in danh_sach_ph_lop:
            ph = phhs.phu_huynh
            # Đếm tin nhắn gửi từ PH này, chưa đọc
            so_tin = TinNhan.objects.filter(
                hoi_thoai__giao_vien=gv_hien_tai, hoi_thoai__phu_huynh=ph,
                nguoi_gui=ph.tai_khoan, da_doc=False
            ).count()
            tong_tin_nhan_chua_doc += so_tin

            danh_sach_ph_kem_tin_nhan.append({
                'id': ph.id,
                'ho_ten': ph.ho_ten,
                'hs_ten': phhs.hoc_sinh.ho_ten,
                'chu_cai': ph.ho_ten[0] if ph.ho_ten else '',
                'so_tin_chua_doc': so_tin
            })

    # ========================================================
    # 1. LỌC MÔN HỌC THÔNG MINH (TỔ HỢP + MÔN CHUNG)
    # ========================================================
    danh_sach_mon_cua_lop = []
    if lop_chinh and lop_chinh.to_hop:
        mon_to_hop = lop_chinh.to_hop.danh_sach_mon.all()
        mon_chung = MonHoc.objects.filter(nhom_mon='Chung')
        danh_sach_mon_cua_lop = (mon_to_hop | mon_chung).distinct()
    else:
        danh_sach_mon_cua_lop = MonHoc.objects.all()

    # ========================================================
    # 2. LỌC HỌC KỲ THEO NĂM HỌC CỦA LỚP CHỦ NHIỆM
    # ========================================================
    danh_sach_hk_cua_lop = []
    if lop_chinh:
        # Lấy năm học của lớp (VD: "2025-2026")
        nam_hoc_hien_tai = lop_chinh.nam_hoc
        # Bốc ra các học kỳ thuộc đúng năm học đó
        danh_sach_hk_cua_lop = HocKy.objects.filter(ten_nam_hoc=nam_hoc_hien_tai).order_by('id')


    # ========================================================
    # ĐÓNG GÓI DỮ LIỆU GỬI RA GIAO DIỆN HTML
    # ========================================================
    context = {
        'gv': gv_hien_tai,
        'lop_chu_nhiem': lop_chu_nhiem,
        'lop_chinh': lop_chinh,
        'danh_sach_hs': danh_sach_hs,
        'so_luong_hs': len(danh_sach_hs),
        'danh_sach_don': danh_sach_don,
        'so_don_cho_duyet': so_don_cho_duyet,
        'danh_sach_ph_lop': danh_sach_ph_lop,
        'so_ph': danh_sach_ph_lop.count() if danh_sach_ph_lop else 0,
        
        # 2 BIẾN QUAN TRỌNG NHẤT ĐỂ ĐỔ LÊN DROPDOWN:
        'danh_sach_mon_cua_lop': danh_sach_mon_cua_lop, 
        'danh_sach_hk_cua_lop': danh_sach_hk_cua_lop, 
        'danh_sach_ph_kem_tin_nhan': danh_sach_ph_kem_tin_nhan,
        'tong_tin_nhan_chua_doc': tong_tin_nhan_chua_doc,
        'tong_vang_cp': tong_vang_cp,
        'tong_vang_kp': tong_vang_kp,
        'chuoi_ty_le': chuoi_ty_le,
        'trung_binh_tuan': trung_binh_tuan,
        'thoi_gian_bieu_do': thoi_gian_bieu_do, # THÊM DÒNG NÀY
        'buoi_da_diem_danh_json': json.dumps(buoi_da_diem_danh, ensure_ascii=False),
    }
    
    return render(request, "teacher_dashboard.html", context)

from django.shortcuts import render
from django.contrib.auth.decorators import login_required
from .models import HocSinh, BangDiem

# ================= GIAO DIỆN HỌC SINH (ĐÃ NÂNG CẤP DỮ LIỆU KÉP HK1 & HK2) =================
@login_required(login_url='login')
def student_dashboard(request):
    if request.user.vai_tro != 'hs':
        return redirect('login')

    try:
        hs = request.user.hoc_sinh
        from .models import BangDiem, DiemDanh, ThongBao
        
        # 1. Kéo TOÀN BỘ bảng điểm của em học sinh này
        tat_ca_diem = BangDiem.objects.filter(hoc_sinh=hs).select_related('mon_hoc', 'hoc_ky')
        
        # 2. Tạo rổ gom nhóm điểm
        bang_diem_gop = {}
        for diem in tat_ca_diem:
            mon_id = diem.mon_hoc.id
            if mon_id not in bang_diem_gop:
                bang_diem_gop[mon_id] = {
                    'mon_hoc': diem.mon_hoc,
                    'hk1': None,
                    'hk2': None,
                    'hanh_kiem': ''
                }
            
            # --- ĐOẠN KHÁNG LỖI TÊN HỌC KỲ ---
            # Chuyển tên học kỳ thành chữ thường để dễ nhận diện
            ten_ky = str(diem.hoc_ky.ky_hoc).lower()
            
            # Nếu tên học kỳ có chứa số '1' (Bắt được cả: HK1, Học kỳ 1, Kỳ 1...)
            if '1' in ten_ky:
                bang_diem_gop[mon_id]['hk1'] = diem
                if diem.hanh_kiem and not bang_diem_gop[mon_id]['hanh_kiem']: 
                    bang_diem_gop[mon_id]['hanh_kiem'] = diem.hanh_kiem
                    
            # Nếu tên học kỳ có chứa số '2' (Bắt được cả: HK2, Học kỳ 2, Kỳ 2...)
            elif '2' in ten_ky:
                bang_diem_gop[mon_id]['hk2'] = diem
                if diem.hanh_kiem: 
                    bang_diem_gop[mon_id]['hanh_kiem'] = diem.hanh_kiem

        # 3. Chuyển rổ thành Danh sách đẩy ra HTML
        danh_sach_diem = bang_diem_gop.values()

        # 4. Thống kê dữ liệu khác
        tong_buoi = DiemDanh.objects.filter(hoc_sinh=hs).count()
        vang_cp = DiemDanh.objects.filter(hoc_sinh=hs, trang_thai='Vang_CP').count()
        vang_kp = DiemDanh.objects.filter(hoc_sinh=hs, trang_thai='Vang_KP').count()
        danh_sach_thong_bao = ThongBao.objects.filter(lop_hoc=hs.lop_hoc).order_by('-ngay_gui')
        # ================= LOGIC TÍNH TOÁN CHUYÊN CẦN =================
        # 1. Tìm 2 học kỳ trong Database (Đã nâng cấp thuật toán tìm kiếm mềm)
        hk1 = HocKy.objects.filter(ky_hoc__icontains='1').last()
        hk2 = HocKy.objects.filter(ky_hoc__icontains='2').last()

        # 2. Hàm gom chung để đếm số liệu điểm danh
        def thong_ke_chuyen_can(hoc_ky_obj):
            # Nếu chưa tạo Học kỳ trong DB thì trả về toàn số 0
            if not hoc_ky_obj:
                return {'tong': 0, 'comat': 0, 'vcp': 0, 'vkp': 0}
                
            # Lấy tất cả các ngày mà GV có bấm lưu điểm danh cho học sinh này TRONG HỌC KỲ ĐÓ
            diem_danh_hk = DiemDanh.objects.filter(
                hoc_sinh=hs, # 'hs' là biến học sinh hiện tại (Bác nhớ check xem biến này tên gì trong hàm của bác nhé)
                ngay__gte=hoc_ky_obj.ngay_bat_dau,
                ngay__lte=hoc_ky_obj.ngay_ket_thuc
            )
            
            return {
                'tong': diem_danh_hk.count(), # Có bao nhiêu dòng là có bấy nhiêu buổi đã điểm danh
                'comat': diem_danh_hk.filter(trang_thai='Co_mat').count(),
                'vcp': diem_danh_hk.filter(trang_thai='Vang_CP').count(),
                'vkp': diem_danh_hk.filter(trang_thai='Vang_KP').count(),
            }

        # 3. Tiến hành đếm
        cc_hk1 = thong_ke_chuyen_can(hk1)
        cc_hk2 = thong_ke_chuyen_can(hk2)

        # 4. Cộng dồn ra dữ liệu Cả năm
        cc_cn = {
            'tong': cc_hk1['tong'] + cc_hk2['tong'],
            'comat': cc_hk1['comat'] + cc_hk2['comat'],
            'vcp': cc_hk1['vcp'] + cc_hk2['vcp'],
            'vkp': cc_hk1['vkp'] + cc_hk2['vkp'],
        }
        # ================= LOGIC LỊCH SỬ CHI TIẾT NGHỈ HỌC =================
        # 1. Lấy tất cả các ngày bị đánh vắng (Có phép & Không phép) từ bảng Điểm Danh
        danh_sach_vang = DiemDanh.objects.filter(
            hoc_sinh=hs,
            trang_thai__in=['Vang_CP', 'Vang_KP']
        ).order_by('-ngay') # Sắp xếp ngày mới nhất lên đầu

        lich_su_chi_tiet = []
        for vang in danh_sach_vang:
            # 2. Tìm xem trong ngày vắng này, phụ huynh có làm đơn xin nghỉ không?
            don_xin = XinNghiPhep.objects.filter(
                hoc_sinh=hs,
                ngay_bat_dau__lte=vang.ngay,
                ngay_ket_thuc__gte=vang.ngay
            ).first()

            # 3. Chắp vá dữ liệu từ 2 nguồn (Giáo viên + Phụ huynh)
            if don_xin:
                ly_do = don_xin.ly_do
                # Trường hợp 1: Đã Duyệt
                if don_xin.trang_thai == 'Da_Duyet':
                    ghi_chu = "GVCN đã xem đơn và CHẤP THUẬN cho nghỉ phép"
                
                # Trường hợp 2: Bị Từ chối
                elif don_xin.trang_thai == 'Tu_Choi':
                    # Cô từ chối đơn, và khi điểm danh đánh luôn là Vắng Không Phép
                    if vang.trang_thai == 'Vang_KP':
                        ghi_chu = "Đơn xin nghỉ BỊ TỪ CHỐI, GV đánh vắng không phép"
                    # Trường hợp hi hữu: Từ chối đơn nhưng lúc điểm danh lại lỡ tay đánh Có phép
                    else:
                        ghi_chu = "Đơn xin nghỉ bị từ chối, nhưng GV châm chước ghi nhận có phép"
                
                # Trường hợp 3: Đang chờ duyệt (Phụ huynh nộp trễ sau khi cô đã điểm danh)
                else: 
                    if vang.trang_thai == 'Vang_KP':
                        ghi_chu = "Phụ huynh nộp đơn trễ, GV đã lỡ điểm danh vắng không phép (Chờ duyệt)"
                    else:
                        ghi_chu = "Đơn vừa được nộp, đang chờ GVCN duyệt chính thức"
            else:
                ly_do = "Chưa rõ lý do"
                # Cô đánh có phép nhưng không thấy đơn trên hệ thống
                if vang.trang_thai == 'Vang_CP':
                    ghi_chu = "GV ghi nhận có phép (Do Phụ huynh gọi điện/nhắn tin trực tiếp)"
                # Cô đánh không phép và cũng không thấy đơn
                else:
                    ghi_chu = "Học sinh vắng mặt không lý do, đang chờ Phụ huynh giải trình"

            # Đóng gói dữ liệu
            lich_su_chi_tiet.append({
                'ngay': vang.ngay,
                'loai': vang.trang_thai,
                'ly_do': ly_do,
                'ghi_chu': ghi_chu
            })
        context = {
            'hs': hs,
            'cc_hk1': cc_hk1, 
            'cc_hk2': cc_hk2,
            'cc_cn': cc_cn,
            'lich_su_chi_tiet': lich_su_chi_tiet,
            'danh_sach_diem': danh_sach_diem,
            'tong_buoi': tong_buoi,
            'vang_cp': vang_cp,
            'vang_kp': vang_kp,
            'danh_sach_thong_bao': danh_sach_thong_bao,
        }
        return render(request, "student_dashboard.html", context)
    except Exception as e:
        messages.error(request, f"Lỗi tải dữ liệu học sinh! {e}")
        return redirect('login')
# ================= GIAO DIỆN PHỤ HUYNH =================
@login_required(login_url='login')
def parent_dashboard(request):
    if request.user.vai_tro != 'ph':
        return redirect('login')

    try:
        ph = request.user.phu_huynh
        # Nhớ thêm TinNhan vào đây để đếm
        from .models import PhuHuynhHocSinh, BangDiem, DiemDanh, ThongBao, XinNghiPhep, TinNhan
        
        danh_sach_con_qs = PhuHuynhHocSinh.objects.filter(phu_huynh=ph).select_related('hoc_sinh__lop_hoc')
        danh_sach_con = [item.hoc_sinh for item in danh_sach_con_qs]
        
        hs_dang_chon = None
        danh_sach_diem = []
        tong_buoi = vang_cp = vang_kp = 0
        danh_sach_thong_bao = []
        
        # ==========================================================
        # CHỖ ĐƯỢC THÊM MỚI: TÌM GIÁO VIÊN CHỦ NHIỆM VÀ ĐẾM TIN NHẮN
        # ==========================================================
        danh_sach_gv_kem_tin_nhan = []
        tong_tin_nhan_chua_doc = 0
        
        gv_da_xu_ly = set()
        for hs in danh_sach_con:
            gvcn = hs.lop_hoc.giao_vien_cn if hs.lop_hoc else None
            if gvcn and gvcn.id not in gv_da_xu_ly:
                gv_da_xu_ly.add(gvcn.id)
                # Quét DB xem có bao nhiêu tin nhắn GV gửi mà mình chưa đọc
                so_tin = TinNhan.objects.filter(
                    hoi_thoai__giao_vien=gvcn, hoi_thoai__phu_huynh=ph,
                    nguoi_gui=gvcn.tai_khoan, da_doc=False
                ).count()
                tong_tin_nhan_chua_doc += so_tin
                
                danh_sach_gv_kem_tin_nhan.append({
                    'id': gvcn.id, 'ho_ten': gvcn.ho_ten,
                    'chu_cai': gvcn.ho_ten[0] if gvcn.ho_ten else '',
                    'so_tin_chua_doc': so_tin
                })
        # ==========================================================

        danh_sach_don = []

        if danh_sach_con:
            hs_dang_chon = danh_sach_con[0]
            danh_sach_don = XinNghiPhep.objects.filter(hoc_sinh=hs_dang_chon).order_by('-ngay_tao')

            tat_ca_diem = BangDiem.objects.filter(hoc_sinh=hs_dang_chon).select_related('mon_hoc', 'hoc_ky')
            bang_diem_gop = {}
            for diem in tat_ca_diem:
                mon_id = diem.mon_hoc.id
                if mon_id not in bang_diem_gop:
                    bang_diem_gop[mon_id] = {'mon_hoc': diem.mon_hoc, 'hk1': None, 'hk2': None, 'hanh_kiem': ''}
                ten_ky = str(diem.hoc_ky.ky_hoc).lower()
                if '1' in ten_ky:
                    bang_diem_gop[mon_id]['hk1'] = diem
                    if diem.hanh_kiem and not bang_diem_gop[mon_id]['hanh_kiem']: 
                        bang_diem_gop[mon_id]['hanh_kiem'] = diem.hanh_kiem
                elif '2' in ten_ky:
                    bang_diem_gop[mon_id]['hk2'] = diem
                    if diem.hanh_kiem: 
                        bang_diem_gop[mon_id]['hanh_kiem'] = diem.hanh_kiem
            danh_sach_diem = bang_diem_gop.values()

            tong_buoi = DiemDanh.objects.filter(hoc_sinh=hs_dang_chon).count()
            vang_cp = DiemDanh.objects.filter(hoc_sinh=hs_dang_chon, trang_thai='Vang_CP').count()
            vang_kp = DiemDanh.objects.filter(hoc_sinh=hs_dang_chon, trang_thai='Vang_KP').count()
            if hs_dang_chon.lop_hoc:
                danh_sach_thong_bao = ThongBao.objects.filter(lop_hoc=hs_dang_chon.lop_hoc).order_by('-ngay_gui')

        context = {
            'ph': ph,
            'danh_sach_con': danh_sach_con,
            'hs_dang_chon': hs_dang_chon,
            'danh_sach_diem': danh_sach_diem,
            'tong_buoi': tong_buoi,
            'vang_cp': vang_cp,
            'vang_kp': vang_kp,
            'danh_sach_thong_bao': danh_sach_thong_bao,
            'danh_sach_don': danh_sach_don,
            
            # Thay list giáo viên cũ bằng biến đếm tin nhắn mới
            'danh_sach_gv_kem_tin_nhan': danh_sach_gv_kem_tin_nhan,
            'tong_tin_nhan_chua_doc': tong_tin_nhan_chua_doc,
        }
        return render(request, 'parent_dashboard.html', context)
    except Exception as e:
        messages.error(request, f"Lỗi tải dữ liệu phụ huynh! {e}")
        return redirect('login')
#Thêm giáo viên
def them_giao_vien(request):
    if request.method == 'POST':
        ma    = request.POST.get('gv_ma')
        ten   = request.POST.get('gv_ten')
        sdt   = request.POST.get('gv_sdt')
        email = request.POST.get('gv_email')
        mk    = request.POST.get('gv_mk')

        try:
            tk = TaiKhoan.objects.create_user(username=ma, password=mk, vai_tro='gv', email=email, is_active=True)
            gv = GiaoVien.objects.create(
                tai_khoan=tk,
                ho_ten=ten,
                so_dien_thoai=sdt,
                email=email
            )
            return JsonResponse({
                'success': True,
                'message': f'Đã thêm giáo viên {ten} thành công!',
                'giao_vien': {
                    'id':    gv.id,
                    'ma':    ma,
                    'ten':   ten,
                    'sdt':   sdt or 'Chưa có',
                    'email': email or 'Chưa có',
                    'tk_id': tk.id,
                }
            })
        except Exception as e:
            return JsonResponse({
                'success': False,
                'message': f'Lỗi! Mã GV (Username) có thể đã bị trùng! ({e})'
            })

    return JsonResponse({'success': False, 'message': 'Phương thức không hợp lệ!'})

# Thêm học sinh
def them_hoc_sinh(request):
    if request.method == 'POST':
        ma_hs = request.POST.get('hs_user')
        ten   = request.POST.get('hs_ten')
        id_lop = request.POST.get('hs_lop')
        mk    = request.POST.get('hs_mk')
        ngay_sinh = request.POST.get('hs_ngaysinh')
        gioi_tinh = request.POST.get('hs_gioitinh')
        dia_chi   = request.POST.get('hs_diachi')

        try:
            lop = LopHoc.objects.get(id=id_lop) if id_lop else None
            email = request.POST.get('hs_email', '')  # Thêm dòng này
            # them_hoc_sinh
            tk = TaiKhoan.objects.create_user(username=ma_hs, password=mk, vai_tro='hs', is_active=True)
            hs  = HocSinh.objects.create(
                tai_khoan=tk,
                ma_hoc_sinh=ma_hs,
                ho_ten=ten,
                lop_hoc=lop,
                ngay_sinh=ngay_sinh,
                gioi_tinh=gioi_tinh,
                dia_chi=dia_chi
            )
            return JsonResponse({
                'success': True,
                'message': f'Đã thêm học sinh {ten} vào lớp {lop.ten_lop if lop else "chưa xếp"}!',
                'hoc_sinh': {
                    'id':       hs.id,
                    'ma':       ma_hs,
                    'ten':      ten,
                    'email':    email,  # Sửa dòng này
                    'ten_lop':  lop.ten_lop if lop else 'Chưa xếp lớp',
                    'lop_id':   lop.id if lop else '',
                    'lop_khoi': lop.khoi if lop else '',
                    'tk_id':    tk.id,
                }
            })
        except Exception as e:
            return JsonResponse({
                'success': False,
                'message': f'Lỗi: Mã HS (Username) có thể đã bị trùng! ({e})'
            })

    return JsonResponse({'success': False, 'message': 'Phương thức không hợp lệ!'})


def them_phu_huynh(request):
    if request.method == 'POST':
        sdt   = request.POST.get('ph_sdt')
        ten   = request.POST.get('ph_ten')
        id_hs = request.POST.get('ph_hs')
        mk    = request.POST.get('ph_mk')

        try:
            # them_phu_huynh  
            tk = TaiKhoan.objects.create_user(username=sdt, password=mk, vai_tro='ph', is_active=True)
            ph = PhuHuynh.objects.create(
                tai_khoan=tk,
                ho_ten=ten,
                so_dien_thoai=sdt
            )

            ten_hs = ''
            lop_hs = ''
            if id_hs:
                hs = HocSinh.objects.get(id=id_hs)
                PhuHuynhHocSinh.objects.create(
                    phu_huynh=ph,
                    hoc_sinh=hs,
                    quan_he='Phụ huynh'
                )
                ten_hs = hs.ho_ten
                lop_hs = hs.lop_hoc.ten_lop if hs.lop_hoc else 'Chưa xếp lớp'

            return JsonResponse({
                'success': True,
                'message': f'Đã thêm phụ huynh {ten} thành công!',
                'phu_huynh': {
                    'id':    ph.id,
                    'sdt':   sdt,
                    'ten':   ten,
                    'tk_id': tk.id,
                    'ten_hs': ten_hs,
                    'lop_hs': lop_hs,
                    'hs_id':  id_hs or '',
                }
            })
        except Exception as e:
            print(f"LỖI TẠO PHỤ HUYNH: {e}")
            return JsonResponse({
                'success': False,
                'message': f'Lỗi: Số điện thoại này đã tồn tại! ({e})'
            })

    return JsonResponse({'success': False, 'message': 'Phương thức không hợp lệ!'})
#Thêm học kỳ
def them_hoc_ky(request):
    if request.method == 'POST':
        nam = request.POST.get('hk_nam')
        ky = request.POST.get('hk_ky')
        bd = request.POST.get('hk_bd')
        kt = request.POST.get('hk_kt')

        try:
            # Code mới khớp 100% với models.py của bác
            HocKy.objects.create(
                ten_nam_hoc=nam, 
                ky_hoc=ky,
                ngay_bat_dau=bd,
                ngay_ket_thuc=kt,
                trang_thai='Dang_Hoc' # Mặc định theo models của bác
            )
            return JsonResponse({'success': True, 'message': f"Đã thêm {ky} năm học {nam} thành công!"})
        except Exception as e:
            return JsonResponse({'success': False, 'message': f"Lỗi: Không thể tạo học kỳ! Chi tiết: {e}"})
            
    return JsonResponse({'success': False, 'message': 'Phương thức không hợp lệ!'})


# Phân công chủ nhiệm (Đã nâng cấp Ràng buộc 1 Giáo viên - 1 Lớp)
def phan_cong_chu_nhiem(request):
    if request.method == 'POST':
        id_lop = request.POST.get('pc_lop')
        id_gv = request.POST.get('pc_gv')

        try:
            # 1. Lôi cái Lớp học và Giáo viên được chọn ra
            lop = LopHoc.objects.get(id=id_lop)
            gv = GiaoVien.objects.get(id=id_gv)
            
            # 2. KIỂM TRA RÀNG BUỘC (Radar quét toàn trường)
            # Tìm xem có Lớp nào KHÁC lớp hiện tại mà Thầy/Cô này đang làm chủ nhiệm không?
            lop_khac = LopHoc.objects.filter(giao_vien_cn=gv).exclude(id=id_lop).first()
            
            if lop_khac:
                # Nếu phát hiện Thầy/Cô này đã "có chủ" -> Quăng thông báo lỗi màu đỏ và chặn lại ngay!
                return JsonResponse({'success': False, 'message': f"❌ Ràng buộc: Thầy/Cô {gv.ho_ten} hiện đang là chủ nhiệm của lớp {lop_khac.ten_lop}! Mỗi giáo viên chỉ được chủ nhiệm 1 lớp."})
            else:
                # Nếu an toàn (chưa có lớp nào, hoặc đang thao tác lại chính lớp cũ của mình) -> Cho phép lưu
                lop.giao_vien_cn = gv
                lop.save()
                return JsonResponse({'success': True, 'message': f"✅ Đã phân công Thầy/Cô {gv.ho_ten} chủ nhiệm lớp {lop.ten_lop}!"})
                
        except Exception as e:
            print(f"LỖI PHÂN CÔNG: {e}")
            return JsonResponse({'success': False, 'message': "Lỗi: Không thể phân công, vui lòng kiểm tra lại hệ thống!"})
            
    return JsonResponse({'success': False, 'message': 'Phương thức không hợp lệ!'})

def sua_hoc_ky(request):
    if request.method == 'POST':
        print("====== TOÀN BỘ POST DATA ======")
        print(request.POST)
        hk_id = request.POST.get('hk_id')
        print(f"====== hk_id = '{hk_id}' ======")
        hk_id = request.POST.get('hk_id')
        print(f"======== ID NHẬN ĐƯỢC: '{hk_id}' ========")  # debug
        
        if hk_id and hk_id.strip():  # kiểm tra không rỗng
            try:
                updated = HocKy.objects.filter(id=int(hk_id)).update(
                    ten_nam_hoc=request.POST.get('hk_nam'),
                    ky_hoc=request.POST.get('hk_ky'),
                    ngay_bat_dau=request.POST.get('hk_bd'),
                    ngay_ket_thuc=request.POST.get('hk_kt'),
                    trang_thai=request.POST.get('hk_trangthai')
                )
                if updated:
                    return JsonResponse({'success': True, 'message': "Đã cập nhật Học kỳ thành công!"})
                else:
                    return JsonResponse({'success': False, 'message': "Không tìm thấy Học kỳ cần sửa!"})
            except Exception as e:
                return JsonResponse({'success': False, 'message': f"Lỗi: {e}"})
        else:
            return JsonResponse({'success': False, 'message': "Lỗi: Không nhận được ID của Học kỳ!"})
            
    return JsonResponse({'success': False, 'message': 'Phương thức không hợp lệ!'})
# ================= THÊM HÀM XÓA HỌC KỲ =================
def xoa_hoc_ky(request):
    if request.method == 'POST':
        hk_id = request.POST.get('hk_id')
        try:
            hk = HocKy.objects.get(id=hk_id)
            ten_hk = f"{hk.ky_hoc} ({hk.ten_nam_hoc})"
            hk.delete() # Lệnh này sẽ xóa bay màu học kỳ khỏi CSDL
            return JsonResponse({'success': True, 'message': f"Đã xóa {ten_hk} thành công!"})
        except Exception as e:
            return JsonResponse({'success': False, 'message': f"Lỗi không thể xóa học kỳ: {e}"})
    return JsonResponse({'success': False, 'message': 'Phương thức không hợp lệ!'})

def sua_lop_hoc(request):
    if request.method == 'POST':
        lop_id = request.POST.get('lop_id')
        
        if lop_id:
            try:
                # Xử lý ID của Giáo viên chủ nhiệm
                gv_id = request.POST.get('lop_gvcn')
                if gv_id == "": gv_id = None
                
                # Xử lý ID của Tổ hợp môn (THÊM MỚI)
                tohop_id = request.POST.get('lop_tohop')
                if tohop_id == "": tohop_id = None
                
                # DÙNG LỆNH UPDATE TRỰC TIẾP
                LopHoc.objects.filter(id=lop_id).update(
                    ten_lop=request.POST.get('lop_ten'),
                    khoi=request.POST.get('lop_khoi'),
                    nam_hoc=request.POST.get('lop_nam'),
                    giao_vien_cn_id=gv_id,
                    to_hop_id=tohop_id # LƯU TỔ HỢP VÀO DB
                )
                return JsonResponse({'success': True, 'message': "Đã cập nhật Lớp học thành công!"})
            except Exception as e:
                return JsonResponse({'success': False, 'message': f"Lỗi: Không thể cập nhật Lớp học! {e}"})
        else:
            return JsonResponse({'success': False, 'message': "Lỗi: Không tìm thấy ID của Lớp cần sửa!"})
            
    return JsonResponse({'success': False, 'message': 'Phương thức không hợp lệ!'})

def sua_mon_hoc(request):
    if request.method == 'POST':
        mh_id = request.POST.get('mh_id')
        if mh_id:
            try:
                MonHoc.objects.filter(id=mh_id).update(
                    ma_mon_hoc=request.POST.get('mh_ma'),
                    ten_mon_hoc=request.POST.get('mh_ten'),
                    loai_mon=request.POST.get('mh_loai'),
                    so_tiet_ky=request.POST.get('mh_sotiet'),
                    nhom_mon=request.POST.get('mh_nhom')
                )
                return JsonResponse({'success': True, 'message': "Đã cập nhật môn học thành công!"})
            except Exception as e:
                return JsonResponse({'success': False, 'message': f"Lỗi: {e}"})
        else:
            return JsonResponse({'success': False, 'message': "Không tìm thấy môn học cần sửa!"})
    return JsonResponse({'success': False, 'message': 'Phương thức không hợp lệ!'})
# ================= THÊM HÀM XÓA MÔN HỌC =================
def xoa_mon_hoc(request):
    if request.method == 'POST':
        mh_id = request.POST.get('mh_id')
        try:
            mon = MonHoc.objects.get(id=mh_id)
            ten_mon = mon.ten_mon_hoc
            mon.delete() # Xóa bay màu khỏi CSDL
            return JsonResponse({'success': True, 'message': f"Đã xóa môn {ten_mon} thành công!"})
        except Exception as e:
            return JsonResponse({'success': False, 'message': f"Lỗi không thể xóa môn học: {e}"})
    return JsonResponse({'success': False, 'message': 'Phương thức không hợp lệ!'})

def sua_giao_vien(request):
    if request.method == 'POST':
        gv_id = request.POST.get('gv_id')
        try:
            GiaoVien.objects.filter(id=gv_id).update(
                ho_ten=request.POST.get('gv_ten'),
                so_dien_thoai=request.POST.get('gv_sdt'),
                email=request.POST.get('gv_email')
            )
            return JsonResponse({'success': True, 'message': "Đã cập nhật thông tin giáo viên!"})
        except Exception as e:
            return JsonResponse({'success': False, 'message': f"Lỗi: {e}"})
    return JsonResponse({'success': False, 'message': f"Lỗi: {e}"})


def reset_mk_giao_vien(request):
    if request.method == 'POST':
        gv_id = request.POST.get('gv_id')
        mk_moi = request.POST.get('gv_mk_moi')
        mk_xacnhan = request.POST.get('gv_mk_xacnhan')
        
        # Nếu mật khẩu không khớp, trả về lỗi ngay cho AJAX
        if mk_moi != mk_xacnhan:
            return JsonResponse({'success': False, 'message': 'Mật khẩu xác nhận không khớp!'})
            
        try:
            gv = GiaoVien.objects.get(id=gv_id)
            gv.tai_khoan.set_password(mk_moi)
            gv.tai_khoan.save()
            
            # Đổi thành công, trả về JSON cho giao diện
            return JsonResponse({'success': True, 'message': f'Đã đổi mật khẩu cho {gv.ho_ten} thành công!'})
        except Exception as e:
            return JsonResponse({'success': False, 'message': f'Lỗi hệ thống: {e}'})
            
    return JsonResponse({'success': False, 'message': 'Phương thức không hợp lệ!'})


def khoa_giao_vien(request):
    if request.method == 'POST':
        tk_id = request.POST.get('tk_id')
        action = request.POST.get('action')
        try:
            tk = TaiKhoan.objects.get(id=tk_id)
            tk.is_active = (action == 'unlock')
            tk.save()
            trang_thai = "mở khóa" if action == 'unlock' else "khóa"
            return JsonResponse({'success': True, 'message': f"Đã {trang_thai} tài khoản thành công!"})
        except Exception as e:
            return JsonResponse({'success': False, 'message': f"Lỗi: {e}"})
    return JsonResponse({'success': False, 'message': 'Phương thức không hợp lệ!'})

def sua_hoc_sinh(request):
    if request.method == 'POST':
        hs_id = request.POST.get('hs_id')
        try:
            hs = HocSinh.objects.get(id=hs_id)
            hs.ho_ten = request.POST.get('hs_ten')
            
            # --- ÁO GIÁP CHỐNG LỖI NGÀY SINH RỖNG ---
            ngay_sinh = request.POST.get('hs_ngaysinh')
            hs.ngay_sinh = ngay_sinh if ngay_sinh else None 
            # ----------------------------------------
            
            hs.gioi_tinh = request.POST.get('hs_gioitinh')
            hs.dia_chi = request.POST.get('hs_diachi')
            
            # Cập nhật email trên tài khoản
            email = request.POST.get('hs_email')
            if email:
                hs.tai_khoan.email = email
                hs.tai_khoan.save()
                
            # Chuyển lớp nếu có chọn
            lop_id = request.POST.get('hs_lop')
            if lop_id:
                hs.lop_hoc = LopHoc.objects.get(id=lop_id)
                
            hs.save()
            return JsonResponse({'success': True, 'message': f"Đã cập nhật thông tin học sinh {hs.ho_ten}!"})
        except Exception as e:
            return JsonResponse({'success': False, 'message': f"Lỗi cập nhật: {e}"})
    return JsonResponse({'success': False, 'message': 'Phương thức không hợp lệ!'})

def reset_mk_hoc_sinh(request):
    if request.method == 'POST':
        hs_id = request.POST.get('hs_id')
        mk_moi = request.POST.get('hs_mk_moi')
        mk_xacnhan = request.POST.get('hs_mk_xacnhan')
        if mk_moi != mk_xacnhan:
            return JsonResponse({'success': False, 'message': 'Mật khẩu xác nhận không khớp!'})
        try:
            hs = HocSinh.objects.get(id=hs_id)
            hs.tai_khoan.set_password(mk_moi)
            hs.tai_khoan.save()

            # ✅ THÊM DÒNG NÀY
            from django.contrib.auth import update_session_auth_hash
            update_session_auth_hash(request, request.user)

            return JsonResponse({'success': True, 'message': f"Đã đổi mật khẩu cho {hs.ho_ten}!"})
        except Exception as e:
            return JsonResponse({'success': False, 'message': f"Lỗi: {e}"})
    return JsonResponse({'success': False, 'message': 'Phương thức không hợp lệ!'})


# ================= KHÓA / MỞ KHÓA TÀI KHOẢN HỌC SINH =================
def khoa_hoc_sinh(request):
    if request.method == 'POST':
        tk_id = request.POST.get('tk_id')
        action = request.POST.get('action')
        try:
            tk = TaiKhoan.objects.get(id=tk_id)
            tk.is_active = (action == 'unlock')
            tk.save()
            trang_thai = "mở khóa" if action == 'unlock' else "khóa"
            return JsonResponse({'success': True, 'message': f'Đã {trang_thai} tài khoản học sinh thành công!'})
        except Exception as e:
            return JsonResponse({'success': False, 'message': str(e)})
    return JsonResponse({'success': False, 'message': 'Sai phương thức!'})

def sua_phu_huynh(request):
    if request.method == 'POST':
        ph_id = request.POST.get('ph_id')
        try:
            ph = PhuHuynh.objects.get(id=ph_id)
            ph.ho_ten = request.POST.get('ph_ten')
            ph.save()

            # Thêm liên kết học sinh mới nếu có chọn
            hs_id = request.POST.get('ph_hs')
            if hs_id:
                hs = HocSinh.objects.get(id=hs_id)
                # Kiểm tra chưa liên kết thì mới tạo, tránh trùng
                PhuHuynhHocSinh.objects.get_or_create(
                    phu_huynh=ph,
                    hoc_sinh=hs,
                    defaults={'quan_he': 'Phụ huynh'}
                )
            return JsonResponse({'success': True, 'message': f"Đã cập nhật thông tin phụ huynh {ph.ho_ten}!"})
        except Exception as e:
            print(f"LỖI SỬA PHỤ HUYNH: {e}")
            return JsonResponse({'success': False, 'message': f"Lỗi: {e}"})
    return JsonResponse({'success': False, 'message': 'Phương thức không hợp lệ!'})
#import
def import_lop_hoc(request):
    if request.method == 'POST':
        # Kiểm tra xem người dùng có tải file lên không
        if 'file_excel' not in request.FILES:
            messages.error(request, "Lỗi: Bác chưa chọn file Excel!")
            return redirect('admin_dashboard')

        file = request.FILES['file_excel']

        # Kiểm tra đúng đuôi file Excel chưa
        if not file.name.endswith('.xlsx') and not file.name.endswith('.xls'):
            messages.error(request, "Lỗi: Vui lòng tải lên file định dạng Excel (.xlsx hoặc .xls)!")
            return redirect('admin_dashboard')

        try:
            # Đọc file Excel bằng Pandas
            df = pd.read_excel(file)
            df.columns = df.columns.str.strip()  # ← Xóa khoảng trắng thừa trong tên cột    
            print("=== CÁC CỘT TRONG FILE ===")
            print(list(df.columns))
            print("==========================")
            # Ép tên cột trong Excel phải chuẩn (Để tránh giáo vụ tự chế)
            required_cols = ['Tên lớp', 'Khối', 'Năm học']
            if not all(col in df.columns for col in required_cols):
                messages.error(request, "Lỗi: File Excel phải có đúng 3 cột tiêu đề: 'Tên lớp', 'Khối', 'Năm học'")
                return redirect('admin_dashboard')

            count = 0
            # Quét từng dòng trong Excel
            for index, row in df.iterrows():
                ten = str(row['Tên lớp']).strip()
                khoi = str(row['Khối']).strip()
                nam = str(row['Năm học']).strip()

                # Nếu tên lớp bị trống thì bỏ qua dòng đó
                if pd.isna(row['Tên lớp']) or ten == 'nan' or not ten:
                    continue

                # Tạo mới Lớp học (Dùng get_or_create để lỡ file có lớp trùng thì không bị đẻ thêm)
                lop, created = LopHoc.objects.get_or_create(
                    ten_lop=ten,
                    defaults={'khoi': khoi, 'nam_hoc': nam}
                )
                if created:
                    count += 1

            messages.success(request, f"Tuyệt vời! Đã import thành công {count} lớp học mới từ Excel!")

        except Exception as e:
            print(f"LỖI IMPORT EXCEL: {e}")
            messages.error(request, f"Có lỗi xảy ra khi đọc file: {e}")

    return redirect('admin_dashboard')

# ================= IMPORT GIÁO VIÊN TỪ EXCEL =================
def import_gv_excel(request):
    if request.method == 'POST':
        if 'file_excel' not in request.FILES:
            messages.error(request, "Lỗi: Bác chưa chọn file Excel!")
            return redirect('admin_dashboard')
            
        file = request.FILES['file_excel']
        try:
            df = pd.read_excel(file)
            required_cols = ['Mã GV', 'Họ tên', 'Số điện thoại', 'Email']
            if not all(col in df.columns for col in required_cols):
                messages.error(request, "Lỗi: File Excel phải có đủ 4 cột: 'Mã GV', 'Họ tên', 'Số điện thoại', 'Email'")
                return redirect('admin_dashboard')

            count = 0
            for index, row in df.iterrows():
                ma = str(row['Mã GV']).strip()
                ten = str(row['Họ tên']).strip()
                sdt = str(row['Số điện thoại']).strip().replace('.0', '') # Xóa số .0 nếu Excel tự thêm
                email = str(row['Email']).strip()

                if pd.isna(row['Mã GV']) or ma == 'nan': continue

                # Kiểm tra nếu chưa có tài khoản thì mới tạo
                if not TaiKhoan.objects.filter(username=ma).exists():
                    tk = TaiKhoan.objects.create_user(username=ma, password='123', email=email, vai_tro='gv')
                    GiaoVien.objects.create(tai_khoan=tk, ho_ten=ten, so_dien_thoai=sdt, email=email)
                    count += 1

            messages.success(request, f"Đã import thành công {count} Giáo viên! (Mật khẩu mặc định: 123)")
        except Exception as e:
            messages.error(request, f"Lỗi đọc file: {e}")
    return redirect('admin_dashboard')

# ================= IMPORT HỌC SINH TỪ EXCEL =================
def import_hs_excel(request):
    if request.method == 'POST':
        if 'file_excel' not in request.FILES: return redirect('admin_dashboard')
        file = request.FILES['file_excel']
        try:
            df = pd.read_excel(file)
            required_cols = ['Mã HS', 'Họ tên', 'Email', 'Tên Lớp']
            if not all(col in df.columns for col in required_cols):
                messages.error(request, "Lỗi: File Excel cần 4 cột: 'Mã HS', 'Họ tên', 'Email', 'Tên Lớp'")
                return redirect('admin_dashboard')

            count = 0
            for index, row in df.iterrows():
                ma = str(row['Mã HS']).strip()
                ten = str(row['Họ tên']).strip()
                email = str(row['Email']).strip()
                ten_lop = str(row['Tên Lớp']).strip()

                if pd.isna(row['Mã HS']) or ma == 'nan': continue

                # Tìm Lớp học dựa vào tên lớp trong Excel
                lop_hoc = LopHoc.objects.filter(ten_lop=ten_lop).first()

                if not TaiKhoan.objects.filter(username=ma).exists():
                    tk = TaiKhoan.objects.create_user(username=ma, password='123', email=email, vai_tro='hs')
                    HocSinh.objects.create(tai_khoan=tk, ma_hoc_sinh=ma, ho_ten=ten, lop_hoc=lop_hoc)
                    count += 1

            messages.success(request, f"Đã import {count} Học sinh! (Mật khẩu mặc định: 123)")
        except Exception as e:
            messages.error(request, f"Lỗi đọc file: {e}")
    return redirect('admin_dashboard')

# ================= IMPORT PHỤ HUYNH TỪ EXCEL =================
def import_ph_excel(request):
    if request.method == 'POST':
        if 'file_excel' not in request.FILES: return redirect('admin_dashboard')
        file = request.FILES['file_excel']
        try:
            df = pd.read_excel(file)
            required_cols = ['Số điện thoại', 'Họ tên PH', 'Mã HS liên kết']
            if not all(col in df.columns for col in required_cols):
                messages.error(request, "Lỗi: Cần 3 cột: 'Số điện thoại', 'Họ tên PH', 'Mã HS liên kết'")
                return redirect('admin_dashboard')

            count = 0
            for index, row in df.iterrows():
                sdt = str(row['Số điện thoại']).strip().replace('.0', '')
                ten = str(row['Họ tên PH']).strip()
                ma_hs = str(row['Mã HS liên kết']).strip()

                if pd.isna(row['Số điện thoại']) or sdt == 'nan': continue

                # Tìm học sinh để liên kết
                hs_lien_ket = HocSinh.objects.filter(ma_hoc_sinh=ma_hs).first()

                if not TaiKhoan.objects.filter(username=sdt).exists():
                    tk = TaiKhoan.objects.create_user(username=sdt, password='123', vai_tro='ph')
                    ph = PhuHuynh.objects.create(tai_khoan=tk, ho_ten=ten, so_dien_thoai=sdt)
                    
                    if hs_lien_ket:
                        PhuHuynhHocSinh.objects.create(phu_huynh=ph, hoc_sinh=hs_lien_ket, quan_he="Phụ huynh")
                    count += 1

            messages.success(request, f"Đã import {count} Phụ huynh! (Mật khẩu mặc định: 123)")
        except Exception as e:
            messages.error(request, f"Lỗi đọc file: {e}")
    return redirect('admin_dashboard')

# ================= API: GIÁO VIÊN DUYỆT ĐƠN XIN NGHỈ =================
@login_required(login_url='login')
def duyet_don_nghi_phep(request):
    if request.method == 'POST':
        don_id = request.POST.get('don_id')
        hanh_dong = request.POST.get('hanh_dong')  # Nhận chữ 'duyet' hoặc 'tuchoi'
        try:
            from .models import XinNghiPhep
            don = XinNghiPhep.objects.get(id=don_id)
            
            if hanh_dong == 'duyet':
                don.trang_thai = 'Da_Duyet'
                don.save() # Bơm trạng thái mới vào CSDL
                return JsonResponse({'success': True, 'message': 'Đã phê duyệt đơn xin nghỉ!'})
                
            elif hanh_dong == 'tuchoi':
                don.trang_thai = 'Tu_Choi'
                don.save() # Bơm trạng thái mới vào CSDL
                return JsonResponse({'success': True, 'message': 'Đã từ chối đơn xin nghỉ!'})
                
        except Exception as e:
            return JsonResponse({'success': False, 'message': f'Lỗi hệ thống: {str(e)}'})
            
    return JsonResponse({'success': False, 'message': 'Sai phương thức truy cập!'})

# ================= LƯU ĐIỂM DANH VÀO DATABASE CSDL =================
@login_required(login_url='login')
def luu_diem_danh(request):
    if request.method == 'POST':
        try:
            # Import các bảng cần dùng cho cả Điểm danh và Tin nhắn
            from .models import DiemDanh, PhuHuynhHocSinh, HoiThoai, TinNhan
            
            data = json.loads(request.body)
            ngay_diem_danh = data.get('ngay')       
            buoi_hoc = data.get('buoi_hoc', 'Sáng')         
            danh_sach = data.get('danh_sach', [])   
            
            # Xử lý ngày tháng chuẩn hóa để CSDL không báo lỗi
            try:
                ngay_sach = re.sub(r'[^\d/]', '', str(ngay_diem_danh))
                ngay_chuan = datetime.strptime(ngay_sach, '%d/%m/%Y').strftime('%Y-%m-%d')
            except Exception:
                ngay_chuan = datetime.now().strftime('%Y-%m-%d')
            
            # Lấy thông tin tài khoản và hồ sơ Giáo viên đang thao tác
            user_gv = request.user
            gv_hien_tai = user_gv.giao_vien
            
            # Bơm vào Database
            for item in danh_sach:
                hs_obj = HocSinh.objects.get(id=item['hs_id'])
                trang_thai_dd = item['trang_thai']
                
                # 1. Lưu trạng thái điểm danh vào CSDL
                DiemDanh.objects.update_or_create(
                    hoc_sinh=hs_obj,
                    ngay=ngay_chuan,
                    buoi_hoc=buoi_hoc,
                    defaults={
                        'giao_vien': gv_hien_tai,
                        'trang_thai': trang_thai_dd
                    }
                )

                # ================================================================
                # 2. ĐỘNG CƠ TỰ ĐỘNG NHẮN TIN CHO PHỤ HUYNH NẾU VẮNG KHÔNG PHÉP
                # ================================================================
                if trang_thai_dd == 'Vang_KP':
                    # Tìm tất cả phụ huynh đang liên kết với học sinh này
                    danh_sach_phhs = PhuHuynhHocSinh.objects.filter(hoc_sinh=hs_obj)
                    
                    for phhs in danh_sach_phhs:
                        ph = phhs.phu_huynh
                        
                        # Kết nối đường dây liên lạc (Hộp thoại) giữa GV và PH này
                        hoi_thoai, _ = HoiThoai.objects.get_or_create(giao_vien=gv_hien_tai, phu_huynh=ph)
                        
                        # Lên kịch bản nội dung tin nhắn tự động
                        noi_dung_robot = f"HỆ THỐNG THÔNG BÁO TỰ ĐỘNG:\nGiáo viên ghi nhận em {hs_obj.ho_ten} vắng mặt KHÔNG PHÉP buổi {buoi_hoc}, ngày {ngay_diem_danh}. Kính mong phụ huynh kiểm tra và phản hồi lại cho giáo viên."
                        
                        # Bắn tin nhắn vào CSDL
                        TinNhan.objects.create(
                            hoi_thoai=hoi_thoai,
                            nguoi_gui=user_gv,  # Tính là GV gửi
                            noi_dung=noi_dung_robot,
                            da_doc=False
                        )
                        # Cập nhật thời gian hộp thoại để nó nhảy lên đầu danh bạ
                        hoi_thoai.save()
                # ================================================================

            return JsonResponse({'success': True, 'message': 'Đã lưu điểm danh và gửi tin nhắn tự động (nếu có)!'})
            
        except Exception as e:
            return JsonResponse({'success': False, 'message': str(e)})
            
    return JsonResponse({'success': False, 'message': 'Chỉ nhận POST'})

# ================= AJAX: XEM LỊCH SỬ ĐIỂM DANH (ĐỌC TỪ CSDL) =================
@login_required(login_url='login')
def api_lich_su_diem_danh(request):
    try:
        gv = request.user.giao_vien
        lop = gv.lop_chu_nhiem.first()
        
        # Lấy toàn bộ điểm danh của lớp này, sắp xếp ngày mới nhất lên đầu
        diem_danh_list = DiemDanh.objects.filter(hoc_sinh__lop_hoc=lop).order_by('-ngay', '-buoi_hoc')
        
        hist_dict = {}
        for dd in diem_danh_list:
            key = f"{dd.ngay}_{dd.buoi_hoc}"
            if key not in hist_dict:
                hist_dict[key] = {
                    'ngay': dd.ngay.strftime('%d/%m/%Y'),
                    'buoi': dd.buoi_hoc,
                    'co_mat': 0, 'vang_cp': 0, 'vang_kp': 0, 'total': 0
                }
            hist_dict[key]['total'] += 1
            if dd.trang_thai == 'Co_mat': hist_dict[key]['co_mat'] += 1
            elif dd.trang_thai == 'Vang_CP': hist_dict[key]['vang_cp'] += 1
            elif dd.trang_thai == 'Vang_KP': hist_dict[key]['vang_kp'] += 1
            
        return JsonResponse({'success': True, 'data': list(hist_dict.values())})
    except Exception as e:
        return JsonResponse({'success': False, 'message': str(e)})

# ================= AJAX: XEM TỔNG KẾT ĐIỂM DANH TỪNG HS (ĐỌC TỪ CSDL) =================
@login_required(login_url='login')
def api_tong_ket_diem_danh(request):
    try:
        gv = request.user.giao_vien
        lop = gv.lop_chu_nhiem.first()
        data = []
        
        # Quét từng học sinh, đếm số buổi vắng trong CSDL
        for hs in lop.danh_sach_hs.all().order_by('ho_ten'):
            v_cp = DiemDanh.objects.filter(hoc_sinh=hs, trang_thai='Vang_CP').count()
            v_kp = DiemDanh.objects.filter(hoc_sinh=hs, trang_thai='Vang_KP').count()
            data.append({
                'ten': hs.ho_ten,
                'v_cp': v_cp,
                'v_kp': v_kp
            })
            
        return JsonResponse({'success': True, 'data': data})
    except Exception as e:
        return JsonResponse({'success': False, 'message': str(e)})
    
# ================= CÁC HÀM THÊM DỮ LIỆU ĐÃ CHUẨN HÓA =================

def them_mon_hoc(request):
    if request.method == 'POST':
        ma = request.POST.get('mh_ma')
        ten = request.POST.get('mh_ten')
        loai = request.POST.get('mh_loai')
        sotiet = request.POST.get('mh_sotiet')
        nhom = request.POST.get('mh_nhom') 

        try:
            MonHoc.objects.create(
                ma_mon_hoc=ma, ten_mon_hoc=ten, loai_mon=loai, 
                so_tiet_ky=sotiet, nhom_mon=nhom 
            )
            return JsonResponse({'success': True, 'message': f"Đã thêm môn {ten} vào chương trình!"})
        except Exception as e:
            return JsonResponse({'success': True, 'message': f"Đã thêm môn {ten} vào chương trình!"})
    return JsonResponse({'success': False, 'message': 'Phương thức không hợp lệ!'})


def them_lop_hoc(request):
    if request.method == 'POST':
        ten = request.POST.get('lop_ten')
        khoi = request.POST.get('lop_khoi')
        nam = request.POST.get('lop_nam')
        id_tohop = request.POST.get('lop_tohop') 
        id_gvcn = request.POST.get('lop_gvcn') # Đã khôi phục nhận diện Giáo viên chủ nhiệm

        try:
            to_hop_obj = ToHop.objects.get(id=id_tohop) if id_tohop else None
            gvcn_obj = GiaoVien.objects.get(id=id_gvcn) if id_gvcn else None
            
            LopHoc.objects.create(
                ten_lop=ten, 
                khoi=khoi, 
                nam_hoc=nam, 
                to_hop=to_hop_obj,
                giao_vien_cn=gvcn_obj
            )
            return JsonResponse({'success': True, 'message': f"Đã tạo lớp {ten} thành công!"})
        except Exception as e:
            return JsonResponse({'success': False, 'message': "Có lỗi xảy ra khi tạo lớp!"})
    return JsonResponse({'success': False, 'message': 'Phương thức không hợp lệ!'})


# ĐÂY LÀ HÀM TẠO TỔ HỢP MỚI
def them_to_hop(request):
    if request.method == 'POST':
        ten_to_hop = request.POST.get('th_ten')
        danh_sach_mon_ids = request.POST.getlist('th_monhoc') # Lấy danh sách ID môn học đánh dấu tick
        try:
            th = ToHop.objects.create(ten_to_hop=ten_to_hop)
            if danh_sach_mon_ids:
                mon_hocs = MonHoc.objects.filter(id__in=danh_sach_mon_ids)
                th.danh_sach_mon.set(mon_hocs) # Lưu chuẩn ManyToMany
            return JsonResponse({'success': True, 'message': f"Đã tạo {ten_to_hop} thành công với {len(danh_sach_mon_ids)} môn!"})
        except Exception as e:
            return JsonResponse({'success': False, 'message': f"Lỗi tạo tổ hợp: {str(e)}"})
    return JsonResponse({'success': False, 'message': f"Lỗi tạo tổ hợp: {str(e)}"})
# ================= IMPORT ĐIỂM SỐ TỪ EXCEL =================
@login_required(login_url='login')
def import_diem_excel(request):
    if request.method == 'POST':
        mon_id = request.POST.get('mon_hoc_id')
        hk_id = request.POST.get('hoc_ky_id')
        file = request.FILES.get('file_excel')

        if not file or not mon_id or not hk_id:
            messages.error(request, "Lỗi: Thiếu thông tin Môn học, Học kỳ hoặc File Excel!")
            return redirect('teacher_dashboard')

        try:
            from .models import BangDiem, MonHoc, HocKy, HocSinh 
            mon_obj = MonHoc.objects.get(id=mon_id)
            hk_obj = HocKy.objects.get(id=hk_id)
            gv_obj = request.user.giao_vien

            df = pd.read_excel(file)
            header_idx = None
            for i, row in df.iterrows():
                if 'Mã HS' in [str(v).strip() for v in row.values]:
                    header_idx = i
                    break
            
            if header_idx is not None:
                df.columns = df.iloc[header_idx]
                df = df[header_idx + 1:]
            
            df.columns = df.columns.astype(str).str.strip()

            required_cols = ['Mã HS', 'ĐTX 1', 'ĐTX 2', 'ĐTX 3', 'ĐGK', 'ĐCK', 'Hạnh kiểm']
            if not all(col in df.columns for col in required_cols):
                messages.error(request, "Lỗi: Cấu trúc File Excel không đúng chuẩn!")
                return redirect('teacher_dashboard')

            def to_float(val):
                if pd.isna(val) or str(val).strip() == '' or str(val).strip() == 'nan': return None
                try: return float(val)
                except: return None
                
            def khac_nhau(v1, v2):
                s1 = "{:.1f}".format(v1) if v1 is not None else ""
                s2 = "{:.1f}".format(v2) if v2 is not None else ""
                return s1 != s2

            diem_cu_qs = BangDiem.objects.filter(mon_hoc=mon_obj, hoc_ky=hk_obj)
            diem_cu_dict = {bd.hoc_sinh.ma_hoc_sinh: bd for bd in diem_cu_qs} 

            count_thanh_cong = 0
            hs_ids_duoc_nhap = [] 

            for index, row in df.iterrows():
                ma_hs = str(row['Mã HS']).strip().replace('.0', '')
                if not ma_hs or ma_hs == 'nan': continue
                
                hs_obj = HocSinh.objects.filter(ma_hoc_sinh=ma_hs).first()
                if not hs_obj: continue

                hanh_kiem = str(row.get('Hạnh kiểm', '')).strip()
                if pd.isna(row.get('Hạnh kiểm')) or hanh_kiem == 'nan': hanh_kiem = ''

                new_tx1 = to_float(row.get('ĐTX 1'))
                new_tx2 = to_float(row.get('ĐTX 2'))
                new_tx3 = to_float(row.get('ĐTX 3'))
                new_gk  = to_float(row.get('ĐGK'))
                new_ck  = to_float(row.get('ĐCK'))

                bd_cu = diem_cu_dict.get(ma_hs)
                old_tx1 = to_float(bd_cu.diem_tx1) if bd_cu else None
                old_tx2 = to_float(bd_cu.diem_tx2) if bd_cu else None
                old_tx3 = to_float(bd_cu.diem_tx3) if bd_cu else None
                old_gk = to_float(bd_cu.diem_gk) if bd_cu else None
                old_ck = to_float(bd_cu.diem_ck) if bd_cu else None
                old_hk = str(bd_cu.hanh_kiem).strip() if (bd_cu and bd_cu.hanh_kiem) else ''

                co_thay_doi = False
                if khac_nhau(new_tx1, old_tx1): co_thay_doi = True
                if khac_nhau(new_tx2, old_tx2): co_thay_doi = True
                if khac_nhau(new_tx3, old_tx3): co_thay_doi = True
                if khac_nhau(new_gk, old_gk): co_thay_doi = True
                if khac_nhau(new_ck, old_ck): co_thay_doi = True
                if hanh_kiem != old_hk: co_thay_doi = True

                if co_thay_doi:
                    hs_ids_duoc_nhap.append(hs_obj.id)

                BangDiem.objects.update_or_create(
                    hoc_sinh=hs_obj, mon_hoc=mon_obj, hoc_ky=hk_obj,
                    defaults={'giao_vien_nhap': gv_obj, 'diem_tx1': new_tx1, 'diem_tx2': new_tx2, 'diem_tx3': new_tx3, 'diem_gk': new_gk, 'diem_ck': new_ck, 'hanh_kiem': hanh_kiem}
                )
                count_thanh_cong += 1
                
                if hs_obj.lop_hoc and hs_obj.lop_hoc.giao_vien_cn == gv_obj:
                    BangDiem.objects.filter(hoc_sinh=hs_obj, hoc_ky=hk_obj).update(hanh_kiem=hanh_kiem)

# =========================================================
            # 4. ĐỘNG CƠ IMPORT EXCEL: BÁO CHO TẤT CẢ
            # =========================================================
            lop_cn = gv_obj.lop_chu_nhiem.first()
            if lop_cn and len(hs_ids_duoc_nhap) > 0:
                from .models import ThongBao, PhuHuynhHocSinh, HoiThoai, TinNhan

                # LUỒNG 1: Đăng 1 tin lên Bảng thông báo chung của lớp (Không điều kiện phần trăm)
                ThongBao.objects.create(
                    giao_vien=gv_obj, lop_hoc=lop_cn,
                    tieu_de=f"Có điểm mới môn {mon_obj.ten_mon_hoc}",
                    noi_dung=f"Giáo viên vừa đồng bộ bảng điểm từ hệ thống Excel cho môn {mon_obj.ten_mon_hoc} - {hk_obj.ky_hoc}. Các em vào Kết quả học tập kiểm tra nhé!",
                    loai_thong_bao='Thong_Thuong'
                )
                
                # LUỒNG 2: Nhắn tin riêng cho các Phụ huynh có con được cập nhật điểm
                danh_sach_phhs = PhuHuynhHocSinh.objects.filter(hoc_sinh_id__in=hs_ids_duoc_nhap).select_related('hoc_sinh', 'phu_huynh')
                for phhs in danh_sach_phhs:
                    ph = phhs.phu_huynh
                    hs_duoc_nhap = phhs.hoc_sinh
                    hoi_thoai, _ = HoiThoai.objects.get_or_create(giao_vien=gv_obj, phu_huynh=ph)
                    
                    noi_dung_ph = f" HỆ THỐNG: Giáo viên vừa đồng bộ bảng điểm Excel môn {mon_obj.ten_mon_hoc} cho em {hs_duoc_nhap.ho_ten}. Kính mời Phụ huynh kiểm tra chi tiết."
                    TinNhan.objects.create(hoi_thoai=hoi_thoai, nguoi_gui=request.user, noi_dung=noi_dung_ph, da_doc=False)
                    hoi_thoai.save()

            messages.success(request, f"Đã quét qua {count_thanh_cong} dòng hợp lệ. Ghi nhận {len(hs_ids_duoc_nhap)} học sinh có cập nhật điểm!")
        except Exception as e:
            messages.error(request, f"Lỗi đọc file: Vui lòng kiểm tra lại định dạng dữ liệu.")
            
    return redirect('teacher_dashboard')
# ================= ĐỔI MẬT KHẨU PHỤ HUYNH =================
def reset_mk_phu_huynh(request):
    if request.method == 'POST':
        ph_id = request.POST.get('ph_id')
        mk_moi = request.POST.get('ph_mk_moi')
        mk_xacnhan = request.POST.get('ph_mk_xacnhan')
        
        if mk_moi != mk_xacnhan:
            return JsonResponse({'success': False, 'message': 'Mật khẩu xác nhận không khớp!'})
            
        try:
            # Tìm phụ huynh và cập nhật mật khẩu cho tài khoản liên kết
            ph = PhuHuynh.objects.get(id=ph_id)
            ph.tai_khoan.set_password(mk_moi)
            ph.tai_khoan.save()

            return JsonResponse({'success': True, 'message': f"Đã đổi mật khẩu cho phụ huynh {ph.ho_ten} thành công!"})
        except Exception as e:
            return JsonResponse({'success': False, 'message': f"Lỗi: {e}"})       
    return JsonResponse({'success': False, 'message': 'Phương thức không hợp lệ!'})

# ================= KHÓA / MỞ KHÓA TÀI KHOẢN PHỤ HUYNH =================
def khoa_phu_huynh(request):
    if request.method == 'POST':
        tk_id = request.POST.get('tk_id')
        action = request.POST.get('action') # Nhận giá trị 'lock' hoặc 'unlock'
        
        try:
            tk = TaiKhoan.objects.get(id=tk_id)
            # Nếu action là 'unlock' thì is_active = True, ngược lại là False
            tk.is_active = (action == 'unlock')
            tk.save()
            
            trang_thai = "mở khóa" if action == 'unlock' else "khóa"
            return JsonResponse({'success': True, 'message': f"Đã {trang_thai} tài khoản phụ huynh thành công!"})
        except Exception as e:
            return JsonResponse({'success': False, 'message': f"Lỗi: {e}"})
    return JsonResponse({'success': False, 'message': 'Phương thức không hợp lệ!'})
# ================= KÉO ĐIỂM TỪ DATABASE LÊN GIAO DIỆN (BẢN KHÁNG LỖI) =================
@login_required(login_url='login')
def api_lay_diem_lop(request):
    mon_id = request.GET.get('mon_id')
    hk_id = request.GET.get('hk_id')
    
    if not mon_id or not hk_id:
        return JsonResponse({'success': False, 'message': 'Chưa chọn Môn hoặc Học kỳ'})
        
    try:
        from .models import BangDiem
        import math # Thêm thư viện toán học để bắt lỗi NaN
        
        gv = request.user.giao_vien
        lop = gv.lop_chu_nhiem.first()
        
        diem_list = BangDiem.objects.filter(
            hoc_sinh__lop_hoc=lop,
            mon_hoc_id=mon_id,
            hoc_ky_id=hk_id
        )
        
        # MÀNG LỌC: Nếu dính NaN (Not a Number) thì biến thành rỗng để JS khỏi bị Crash
        def sach_se(val):
            if val is None: return ''
            try:
                if math.isnan(float(val)): return ''
                return float(val)
            except:
                return ''

        data = {}
        for d in diem_list:
            data[d.hoc_sinh.id] = {
                'tx1': sach_se(d.diem_tx1),
                'tx2': sach_se(d.diem_tx2),
                'tx3': sach_se(d.diem_tx3),
                'gk': sach_se(d.diem_gk),
                'ck': sach_se(d.diem_ck),
                'hk': d.hanh_kiem if d.hanh_kiem else ''
            }
        return JsonResponse({'success': True, 'data': data})
    except Exception as e:
        return JsonResponse({'success': False, 'message': str(e)})

# ================= LƯU ĐIỂM SỐ NHẬP THỦ CÔNG TỪ GIAO DIỆN =================
@login_required(login_url='login')
def api_luu_diem_thu_cong(request):
    if request.method == 'POST':
        try:
            import json
            data = json.loads(request.body)
            mon_id = data.get('mon_id')
            hk_id = data.get('hk_id')
            diem_list = data.get('diem_list', [])

            if not mon_id or not hk_id:
                return JsonResponse({'success': False, 'message': 'Vui lòng chọn Môn và Học kỳ!'})

            from .models import BangDiem, MonHoc, HocKy, HocSinh
            mon_obj = MonHoc.objects.get(id=mon_id)
            hk_obj = HocKy.objects.get(id=hk_id)
            gv_obj = request.user.giao_vien

            def to_float(val):
                if val is None or str(val).strip() == '' or str(val).strip() == 'nan': return None
                try: return float(val)
                except: return None
                
            # TUYỆT CHIÊU MỚI: Format về chuỗi 1 số lẻ để so sánh (Khắc phục triệt để 8.1 != 8.0999)
            def khac_nhau(v1, v2):
                s1 = "{:.1f}".format(v1) if v1 is not None else ""
                s2 = "{:.1f}".format(v2) if v2 is not None else ""
                return s1 != s2

            # ========================================================
            # 1. BỘ NHỚ ĐỆM: LẤY ĐIỂM CŨ TỪ DATABASE RA ĐỂ SO SÁNH
            # ========================================================
            diem_cu_qs = BangDiem.objects.filter(mon_hoc=mon_obj, hoc_ky=hk_obj)
            diem_cu_dict = {bd.hoc_sinh_id: bd for bd in diem_cu_qs}

            cot_da_nhap = set()
            hs_ids_duoc_nhap = []

            for d in diem_list:
                hs_id_str = d.get('hs_id')
                if not hs_id_str: continue
                hs_id = int(hs_id_str)
                
                hs_obj = HocSinh.objects.filter(id=hs_id).first()
                if not hs_obj: continue

                new_tx1 = to_float(d.get('tx1'))
                new_tx2 = to_float(d.get('tx2'))
                new_tx3 = to_float(d.get('tx3'))
                new_gk = to_float(d.get('gk'))
                new_ck = to_float(d.get('ck'))
                new_hk = str(d.get('hk', '')).strip()
                
                bd_cu = diem_cu_dict.get(hs_id)
                old_tx1 = to_float(bd_cu.diem_tx1) if bd_cu else None
                old_tx2 = to_float(bd_cu.diem_tx2) if bd_cu else None
                old_tx3 = to_float(bd_cu.diem_tx3) if bd_cu else None
                old_gk = to_float(bd_cu.diem_gk) if bd_cu else None
                old_ck = to_float(bd_cu.diem_ck) if bd_cu else None
                old_hk = str(bd_cu.hanh_kiem).strip() if (bd_cu and bd_cu.hanh_kiem) else ''
                
                # [THÊM MỚI] 1. CÔNG THỨC TÍNH ĐIỂM TRUNG BÌNH HỌC KỲ
                dtb_ky = None
                tx_list = [x for x in [new_tx1, new_tx2, new_tx3] if x is not None]
                if len(tx_list) > 0 and new_gk is not None and new_ck is not None:
                    tong = sum(tx_list) + (new_gk * 2) + (new_ck * 3)
                    he_so = len(tx_list) + 5
                    dtb_ky = round(tong / he_so, 1)

                # ========================================================
                # 2. ĐỐI CHIẾU THÔNG MINH BẰNG CHUỖI
                # ========================================================
                co_thay_doi = False
                if khac_nhau(new_tx1, old_tx1):
                    co_thay_doi = True
                    if new_tx1 is not None: cot_da_nhap.add('TX 1')
                if khac_nhau(new_tx2, old_tx2):
                    co_thay_doi = True
                    if new_tx2 is not None: cot_da_nhap.add('TX 2')
                if khac_nhau(new_tx3, old_tx3):
                    co_thay_doi = True
                    if new_tx3 is not None: cot_da_nhap.add('TX 3')
                if khac_nhau(new_gk, old_gk):
                    co_thay_doi = True
                    if new_gk is not None: cot_da_nhap.add('Giữa kỳ')
                if khac_nhau(new_ck, old_ck):
                    co_thay_doi = True
                    if new_ck is not None: cot_da_nhap.add('Cuối kỳ')
                if new_hk != old_hk:
                    co_thay_doi = True
                    if new_hk != '': cot_da_nhap.add('Hạnh kiểm')

                if co_thay_doi:
                    hs_ids_duoc_nhap.append(hs_id)
                
                BangDiem.objects.update_or_create(
                    hoc_sinh=hs_obj, mon_hoc=mon_obj, hoc_ky=hk_obj,
                    defaults={'giao_vien_nhap': gv_obj, 'diem_tx1': new_tx1, 'diem_tx2': new_tx2, 'diem_tx3': new_tx3, 'diem_gk': new_gk, 'diem_ck': new_ck, 'hanh_kiem': new_hk}
                )

                if hs_obj.lop_hoc_id and hs_obj.lop_hoc.giao_vien_cn_id == gv_obj.id:
                    BangDiem.objects.filter(hoc_sinh=hs_obj, hoc_ky=hk_obj).update(hanh_kiem=new_hk)

        
                # [THÊM MỚI] 3. LƯU VÀO BẢNG TỔNG KẾT MÔN
                from .models import TongKetMon
                # Lấy năm học từ thông tin lớp học của học sinh
                nam_hoc_hien_tai = hs_obj.lop_hoc.nam_hoc if hs_obj.lop_hoc else hk_obj.ten_nam_hoc
                
                tong_ket, _ = TongKetMon.objects.get_or_create(
                    hoc_sinh=hs_obj, mon_hoc=mon_obj, nam_hoc=nam_hoc_hien_tai
                )
                
                # Kiểm tra xem đang lưu cho Học kỳ 1 hay Học kỳ 2
                if '1' in str(hk_obj.ky_hoc):
                    tong_ket.dtb_hk1 = dtb_ky
                elif '2' in str(hk_obj.ky_hoc):
                    tong_ket.dtb_hk2 = dtb_ky
                    
                # Nếu có đủ cả 2 kỳ thì tính luôn Cả Năm
                if tong_ket.dtb_hk1 is not None and tong_ket.dtb_hk2 is not None:
                    tong_ket.dtb_ca_nam = round((tong_ket.dtb_hk1 + tong_ket.dtb_hk2 * 2) / 3, 1)
                else:
                    tong_ket.dtb_ca_nam = None
                    
                tong_ket.save() # Chốt cứng xuống CSDL
            # =========================================================
            # 3. ĐỘNG CƠ KÉP: GỌI TÊN HỌC SINH TRÊN LOA + NHẮN TIN PHỤ HUYNH
            # =========================================================
            lop_cn = gv_obj.lop_chu_nhiem.first()
            if lop_cn and len(hs_ids_duoc_nhap) > 0:
                chuoi_cot = ", ".join(cot_da_nhap) if cot_da_nhap else "điểm số"
                from .models import ThongBao, PhuHuynhHocSinh, HoiThoai, TinNhan, HocSinh
                
                # --- LUỒNG 1: BÁO CHO HỌC SINH (GỌI TÊN ĐÍCH DANH TRÊN BẢNG CHUNG) ---
                # Lấy tên các học sinh vừa được sửa điểm
                cac_hs_duoc_sua = HocSinh.objects.filter(id__in=hs_ids_duoc_nhap)
                danh_sach_ten = [hs.ho_ten for hs in cac_hs_duoc_sua]
                
                # Nếu sửa ít (1-3 em) thì ghi rõ tên, nếu sửa nhiều thì gom thành số lượng
                if len(danh_sach_ten) <= 3:
                    chuoi_ten = "em: " + ", ".join(danh_sach_ten)
                else:
                    chuoi_ten = f"{len(danh_sach_ten)} học sinh"

                ThongBao.objects.create(
                    giao_vien=gv_obj,
                    lop_hoc=lop_cn,
                    tieu_de=f"Cập nhật điểm môn {mon_obj.ten_mon_hoc}",
                    noi_dung=f"Giáo viên vừa cập nhật ({chuoi_cot}) cho {chuoi_ten}. Các em có tên vui lòng vào Kết quả học tập kiểm tra nhé!",
                    loai_thong_bao='Thong_Thuong'
                )

                # --- LUỒNG 2: CHỈ NHẮN TIN KÍN ĐÍCH DANH CHO PHỤ HUYNH BỊ SỬA ĐIỂM ---
                danh_sach_phhs = PhuHuynhHocSinh.objects.filter(hoc_sinh_id__in=hs_ids_duoc_nhap).select_related('hoc_sinh', 'phu_huynh')
                for phhs in danh_sach_phhs:
                    ph = phhs.phu_huynh
                    hs_duoc_nhap = phhs.hoc_sinh
                    hoi_thoai, _ = HoiThoai.objects.get_or_create(giao_vien=gv_obj, phu_huynh=ph)
                    
                    noi_dung_ph = f"HỆ THỐNG: Giáo viên vừa cập nhật ({chuoi_cot}) môn {mon_obj.ten_mon_hoc} cho em {hs_duoc_nhap.ho_ten}. Kính mời Phụ huynh kiểm tra chi tiết."
                    TinNhan.objects.create(hoi_thoai=hoi_thoai, nguoi_gui=request.user, noi_dung=noi_dung_ph, da_doc=False)
                    hoi_thoai.save()
                    
            return JsonResponse({'success': True, 'message': f'Lưu thành công! (Cập nhật cho {len(hs_ids_duoc_nhap)} học sinh)'})
        except Exception as e:
            return JsonResponse({'success': False, 'message': str(e)})
            
    return JsonResponse({'success': False, 'message': 'Sai phương thức!'})

def api_thong_tin_ph(request):
    hs_id = request.GET.get('hs_id')
    hs = HocSinh.objects.get(id=hs_id)
    ph_list = PhuHuynhHocSinh.objects.filter(hoc_sinh=hs)
    data = [{'id': p.phu_huynh.id, 'ho_ten': p.phu_huynh.ho_ten, 'sdt': p.phu_huynh.so_dien_thoai, 'quan_he': p.quan_he} for p in ph_list]
    return JsonResponse({'success': True, 'ph_list': data})

def gui_thong_bao_lop(request):
    if request.method == 'POST':
        try:
            # Lấy thông tin GV đang đăng nhập
            gv = request.user.giao_vien
            # Tìm lớp mà GV này đang chủ nhiệm
            lop_cn = gv.lop_chu_nhiem.first() 
            
            if not lop_cn:
                return JsonResponse({'success': False, 'message': 'Thầy/Cô không chủ nhiệm lớp nào!'})

            tieu_de = request.POST.get('tb_tieu_de')
            noi_dung = request.POST.get('tb_noi_dung')
            muc_do = request.POST.get('tb_muc_do', 'Thong_Thuong')

            # Lưu vào Database
            ThongBao.objects.create(
                giao_vien=gv,
                lop_hoc=lop_cn,
                tieu_de=tieu_de,
                noi_dung=noi_dung,
                loai_thong_bao=muc_do
            )
            return JsonResponse({'success': True, 'message': 'Đã gửi thông báo đến toàn lớp!'})
        except Exception as e:
            return JsonResponse({'success': False, 'message': f'Lỗi: {str(e)}'})
            
    return JsonResponse({'success': False, 'message': 'Phương thức không hợp lệ!'})

from django.contrib.auth.decorators import login_required
from django.http import JsonResponse

# ================= API: PHỤ HUYNH NỘP ĐƠN XIN NGHỈ (CÓ HỖ TRỢ ẢNH) =================
@login_required(login_url='login')
def api_nop_don_xin_nghi(request):
    if request.method == 'POST':
        try:
            # Dùng request.POST để lấy chữ
            hs_id = request.POST.get('hs_id')
            ngay_bd = request.POST.get('ngay_bd')
            ngay_kt = request.POST.get('ngay_kt')
            ly_do = request.POST.get('ly_do')
            
            # Dùng request.FILES để lấy ảnh minh chứng
            anh_minh_chung = request.FILES.get('minh_chung')
            
            from .models import XinNghiPhep, PhuHuynhHocSinh
            ph = request.user.phu_huynh
            
            hop_le = PhuHuynhHocSinh.objects.filter(phu_huynh=ph, hoc_sinh_id=hs_id).exists()
            if not hop_le:
                return JsonResponse({'success': False, 'message': 'Không có quyền gửi đơn!'})

            # Bơm vào Database
            XinNghiPhep.objects.create(
                hoc_sinh_id=hs_id,
                phu_huynh=ph,
                ngay_bat_dau=ngay_bd,
                ngay_ket_thuc=ngay_kt,
                ly_do=ly_do,
                minh_chung=anh_minh_chung, # Ảnh sẽ tự động bay vào thư mục media/don_xin_nghi/
                trang_thai='Cho_Duyet'
            )
            return JsonResponse({'success': True, 'message': 'Đã gửi đơn đính kèm thành công!'})
        except Exception as e:
            return JsonResponse({'success': False, 'message': str(e)})
    return JsonResponse({'success': False, 'message': 'Sai phương thức truy cập!'})

# ================= API: LẤY LỊCH SỬ TIN NHẮN 2 CHIỀU =================
@login_required(login_url='login')
def api_lay_tin_nhan(request):
    doi_tac_id = str(request.GET.get('doi_tac_id')) # Ép thành chuỗi để xử lý chữ 'bot'
    vai_tro = request.user.vai_tro
    
    try:
        from django.utils.timezone import localtime
        
        # 1. Tìm đúng Hộp thoại
        if vai_tro == 'gv':
            gv = request.user.giao_vien
            ph = PhuHuynh.objects.get(id=int(doi_tac_id))
            hoi_thoai, _ = HoiThoai.objects.get_or_create(giao_vien=gv, phu_huynh=ph, is_bot_chat=False)
        elif vai_tro == 'ph':
            ph = request.user.phu_huynh
            if doi_tac_id == 'bot':
                hoi_thoai, _ = HoiThoai.objects.get_or_create(phu_huynh=ph, giao_vien=None, is_bot_chat=True)
            else:
                gv = GiaoVien.objects.get(id=int(doi_tac_id))
                hoi_thoai, _ = HoiThoai.objects.get_or_create(giao_vien=gv, phu_huynh=ph, is_bot_chat=False)
        else:
            return JsonResponse({'success': False, 'message': 'Không có quyền truy cập!'})

        # Đánh dấu đã đọc (bỏ qua tin nhắn của bot vì bot không cần đánh dấu)
        TinNhan.objects.filter(hoi_thoai=hoi_thoai, da_doc=False).exclude(nguoi_gui=request.user).exclude(is_bot=True).update(da_doc=True)
        tin_nhans = TinNhan.objects.filter(hoi_thoai=hoi_thoai).order_by('thoi_gian_gui')
        
        data = []
        for tn in tin_nhans:
            thoi_gian_vn = localtime(tn.thoi_gian_gui)
            
            # 2. Phân loại bóng chat (Ai nằm trái, ai nằm phải)
            if tn.is_bot:
                la_nguoi_gui = False # Bot luôn nằm bên trái màn hình
            else:
                la_nguoi_gui = (tn.nguoi_gui == request.user) # Nếu là mình gửi thì nằm bên phải

            data.append({
                'la_nguoi_gui': la_nguoi_gui,
                'noi_dung': tn.noi_dung,
                'thoi_gian': thoi_gian_vn.strftime('%H:%M - %d/%m')
            })
        return JsonResponse({'success': True, 'data': data})
    except Exception as e:
        return JsonResponse({'success': False, 'message': str(e)})

def lam_sach_lich_su(text):
    """Xóa các thẻ HTML và mã MENU khỏi lịch sử chat để AI không bị nhiễu."""
    if not text: return ""
    text = re.sub(r'<[^>]+>', '\n', text)
    text = re.sub(r'\[MENU:.*?\]', '', text)
    return text.strip()

@login_required(login_url='login')
def api_gui_tin_nhan(request):
    if request.method == 'POST':
        try:
            data = json.loads(request.body)
            doi_tac_id = str(data.get('doi_tac_id'))
            noi_dung = data.get('noi_dung', '').strip()
            vai_tro = request.user.vai_tro

            # =================================================================
            # 1. LUỒNG XỬ LÝ: PHỤ HUYNH CHAT VỚI BOT AI
            # =================================================================
            if vai_tro == 'ph' and doi_tac_id == 'bot':
                phu_huynh = request.user.phu_huynh
                hoi_thoai, _ = HoiThoai.objects.get_or_create(phu_huynh=phu_huynh, giao_vien=None, is_bot_chat=True)
                
                # Lưu tin nhắn của phụ huynh vào DB
                TinNhan.objects.create(hoi_thoai=hoi_thoai, nguoi_gui=request.user, noi_dung=noi_dung)

                # Khởi tạo công cụ (Đã khóa chặt bảo mật IDOR theo phiên đăng nhập)
                tool_instance = CongCuPhuHuynh(phu_huynh)
                noi_dung_lower = noi_dung.lower()

                # --- LỐI TẮT PYTHON (BYPASS AI ĐỂ TIẾT KIỆM TOKEN) ---
                if "tất cả các môn" in noi_dung_lower or "tat ca cac mon" in noi_dung_lower:
                    bot_reply = "Dạ thưa Phụ huynh, dưới đây là bảng điểm tổng kết của con:\n\n" + tool_instance.lay_diem_hoc_sinh()
                    TinNhan.objects.create(hoi_thoai=hoi_thoai, nguoi_gui=None, is_bot=True, noi_dung=bot_reply)
                    return JsonResponse({'success': True})
                
                if "cách xin nghỉ" in noi_dung_lower or "làm đơn xin" in noi_dung_lower:
                    bot_reply = "Dạ thưa Phụ huynh, để tạo đơn xin nghỉ học, Phụ huynh vui lòng làm theo hướng dẫn sau:\n\n" + tool_instance.tra_mau_don_xin_nghi()
                    TinNhan.objects.create(hoi_thoai=hoi_thoai, nguoi_gui=None, is_bot=True, noi_dung=bot_reply)
                    return JsonResponse({'success': True})

                # --- DỰNG LỊCH SỬ HỘI THOẠI (BỘ NHỚ) ---
                # Truy vấn 10 tin nhắn gần nhất và lật ngược lại cho đúng thứ tự thời gian
                lich_su_db = TinNhan.objects.filter(hoi_thoai=hoi_thoai).order_by('-thoi_gian_gui')[:10]
                contents = []
                for tn in reversed(lich_su_db):
                    text_sach = lam_sach_lich_su(tn.noi_dung)
                    if not text_sach: continue
                    role = "model" if tn.is_bot else "user"
                    contents.append(types.Content(role=role, parts=[types.Part.from_text(text=text_sach)]))

                # --- SYSTEM INSTRUCTION (HỒN CỐT CỦA AI) ---
                hom_nay = datetime.now().strftime("%d/%m/%Y")
                danh_sach_con = ", ".join([p.hoc_sinh.ho_ten for p in tool_instance.danh_sach_con])
                
                sys_instruct = f"""Bạn là Trợ lý AI giáo dục của trường học. Hôm nay là ngày {hom_nay}.
Phụ huynh đang chat với bạn có con tên là: {danh_sach_con}.
QUY TẮC BẮT BUỘC:
1. Luôn xưng hô "Dạ thưa Phụ huynh". Gọi học sinh là "con của Phụ huynh" hoặc gọi thẳng tên. TUYỆT ĐỐI KHÔNG xưng "con tôi".
2. KHÔNG BAO GIỜ dùng Emoji.
3. Khi báo điểm, giữ nguyên các ký hiệu TX, GK, CK, ĐTB. Tuyệt đối không tự tính toán hay bịa điểm, chỉ đọc y nguyên dữ liệu do Tool cung cấp.
4. Nếu phụ huynh dùng các từ lóng: "cúp học", "trốn tiết", "nghỉ ngang", hãy hiểu là họ muốn kiểm tra chuyên cần/điểm danh.
5. Nếu bạn cần chèn hướng dẫn xin nghỉ, bắt buộc phải dùng đúng đoạn HTML form từ Tool, không được tự ý sửa đổi cú pháp.
6. Khi cần điều hướng menu ứng dụng, hãy chèn đúng mã [MENU:...] vào cuối câu như hệ thống cũ."""

                # --- KHAI BÁO CÁC HÀM WRAPPER CHO SDK ---
                # SDK của Google cần đọc được signature (tham số) của hàm, nên ta bọc nó lại
                def bot_lay_diem(ten_goi_nho_hoc_sinh: str = "", ten_mon_hoc: str = "") -> str:
                    """Tra cứu bảng điểm, điểm trung bình các môn học của học sinh."""
                    return tool_instance.lay_diem_hoc_sinh(ten_goi_nho_hoc_sinh, ten_mon_hoc)
                
                def bot_kiem_tra_chuyen_can(ten_goi_nho_hoc_sinh: str = "", ngay_cu_the: str = "") -> str:
                    """Tra cứu số ngày nghỉ học, tình trạng điểm danh, lịch sử cúp học."""
                    return tool_instance.kiem_tra_chuyen_can(ten_goi_nho_hoc_sinh, ngay_cu_the)

                tools_list = [bot_lay_diem, bot_kiem_tra_chuyen_can, tra_cuu_noi_quy_truong]

                # --- GỌI GEMINI VỚI TÍNH NĂNG CHATS (AUTO TOOL CALLING) ---
                client = genai.Client(api_key=os.getenv("GEMINI_API_KEY"))
                
                # Tách tin nhắn cuối cùng (câu hỏi hiện tại) ra khỏi lịch sử
                cau_hoi_hien_tai = contents.pop().parts[0].text
                
                # Khởi tạo phiên Chat với lịch sử đã làm sạch
                chat = client.chats.create(
                    model="gemini-2.5-flash",
                    config=types.GenerateContentConfig(
                        system_instruction=sys_instruct,
                        tools=tools_list,
                        temperature=0.2
                    ),
                    history=contents
                )
                
                try:
                    # Gửi câu hỏi hiện tại. SDK sẽ TỰ ĐỘNG gọi các hàm Python nếu cần và chờ kết quả.
                    response = chat.send_message(cau_hoi_hien_tai)
                    bot_reply = response.text if response.text else "Dạ thưa Phụ huynh, hệ thống đang bận xử lý, xin vui lòng thử lại sau."
                except Exception as api_err:
                    bot_reply = f"Dạ thưa Phụ huynh, hệ thống AI đang gặp sự cố kết nối tạm thời. ({str(api_err)})"

                # Lưu phản hồi của AI
                TinNhan.objects.create(hoi_thoai=hoi_thoai, nguoi_gui=None, is_bot=True, noi_dung=bot_reply)
                return JsonResponse({'success': True})
            
            # --- LUỒNG 2: CHAT VỚI GIÁO VIÊN BÌNH THƯỜNG ---
            if vai_tro == 'gv':
                gv = request.user.giao_vien
                ph = PhuHuynh.objects.get(id=int(doi_tac_id))
                hoi_thoai, _ = HoiThoai.objects.get_or_create(giao_vien=gv, phu_huynh=ph, is_bot_chat=False)
            elif vai_tro == 'ph':
                ph = request.user.phu_huynh
                gv = GiaoVien.objects.get(id=int(doi_tac_id))
                hoi_thoai, _ = HoiThoai.objects.get_or_create(giao_vien=gv, phu_huynh=ph, is_bot_chat=False)
            
            TinNhan.objects.create(hoi_thoai=hoi_thoai, nguoi_gui=request.user, noi_dung=noi_dung)
            hoi_thoai.save() 
            return JsonResponse({'success': True})
            
        except Exception as e:
            print(f"\n[!!!] BÁC ƠI HỆ THỐNG BÁO LỖI RỒI: {str(e)}\n")
            return JsonResponse({'success': False, 'message': str(e)})
            
    return JsonResponse({'success': False})
# ================= XUẤT FILE EXCEL MẪU CHO GIÁO VIÊN BỘ MÔN (BẢN ĐẸP) =================
@login_required(login_url='login')
def xuat_file_mau_excel(request):
    mon_id = request.GET.get('mon_id')
    hk_id = request.GET.get('hk_id')

    if not mon_id or not hk_id:
        messages.error(request, "Vui lòng chọn Môn học và Học kỳ trên màn hình trước khi tải file mẫu!")
        return redirect('teacher_dashboard')

    try:
        from .models import BangDiem, MonHoc, HocKy, HocSinh
        import pandas as pd
        import io
        from django.http import HttpResponse
        # Import bộ công cụ vẽ đồ họa cho Excel
        from openpyxl.styles import Font, Alignment, PatternFill, Border, Side 

        gv = request.user.giao_vien
        lop = gv.lop_chu_nhiem.first()
        if not lop:
            messages.error(request, "Thầy/Cô không chủ nhiệm lớp nào!")
            return redirect('teacher_dashboard')

        mon = MonHoc.objects.get(id=mon_id)
        hk = HocKy.objects.get(id=hk_id)

        danh_sach_hs = lop.danh_sach_hs.all().order_by('ho_ten')
        diem_cu_qs = BangDiem.objects.filter(mon_hoc=mon, hoc_ky=hk, hoc_sinh__in=danh_sach_hs)
        dict_diem = {bd.hoc_sinh_id: bd for bd in diem_cu_qs}

        # 1. TẠO CẤU TRÚC DỮ LIỆU
        data = []
        for hs in danh_sach_hs:
            d = dict_diem.get(hs.id)
            data.append({
                'Mã HS': hs.ma_hoc_sinh,
                'Họ tên': hs.ho_ten, 
                'ĐTX 1': d.diem_tx1 if d and d.diem_tx1 is not None else '',
                'ĐTX 2': d.diem_tx2 if d and d.diem_tx2 is not None else '',
                'ĐTX 3': d.diem_tx3 if d and d.diem_tx3 is not None else '',
                'ĐGK': d.diem_gk if d and d.diem_gk is not None else '',
                'ĐCK': d.diem_ck if d and d.diem_ck is not None else '',
                'Hạnh kiểm': d.hanh_kiem if d and d.hanh_kiem else ''
            })

        df = pd.DataFrame(data)
        output = io.BytesIO()

        # 2. VẼ EXCEL XỊN XÒ BẰNG OPENPYXL
        with pd.ExcelWriter(output, engine='openpyxl') as writer:
            # Đẩy bảng dữ liệu xuống dòng số 6 (startrow=5 do đếm từ 0), chừa 5 dòng đầu làm Tiêu đề
            df.to_excel(writer, index=False, sheet_name='BangDiem', startrow=5)
            
            # Lấy quyền điều khiển cái Sheet vừa tạo
            worksheet = writer.sheets['BangDiem']

            # --- KHU VỰC ĐỊNH DẠNG FONT & MÀU SẮC ---
            title_font = Font(name='Arial', size=16, bold=True, color='000000')
            bold_font = Font(name='Arial', size=12, bold=True)
            normal_font = Font(name='Arial', size=11)
            note_font = Font(name='Arial', size=10, italic=True, color='FF0000') # Chữ đỏ in nghiêng
            center_align = Alignment(horizontal='center', vertical='center')
            left_align = Alignment(horizontal='left', vertical='center')
            
            # Khung viền mỏng cho các ô
            thin_border = Border(left=Side(style='thin'), right=Side(style='thin'), top=Side(style='thin'), bottom=Side(style='thin'))
            # Màu nền xanh nhạt cho thẻ Tiêu đề bảng
            header_fill = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")

            # --- VẼ TIÊU ĐỀ (HEADER) TRÊN CÙNG ---
            # Dòng 1: TÊN BẢNG ĐIỂM
            worksheet.merge_cells('A1:H1')
            cell = worksheet['A1']
            cell.value = f"BẢNG ĐIỂM BỘ MÔN {mon.ten_mon_hoc.upper()}"
            cell.font = title_font
            cell.alignment = center_align

            # Dòng 2: THÔNG TIN LỚP & HỌC KỲ
            worksheet.merge_cells('A2:H2')
            cell = worksheet['A2']
            cell.value = f"Lớp: {lop.ten_lop}  |  {hk.ky_hoc} - Năm học: {hk.ten_nam_hoc}"
            cell.font = bold_font
            cell.alignment = center_align

            # Dòng 3: GIÁO VIÊN CHỦ NHIỆM
            worksheet.merge_cells('A3:H3')
            cell = worksheet['A3']
            cell.value = f"Giáo viên chủ nhiệm: {gv.ho_ten}"
            cell.font = normal_font
            cell.alignment = center_align

            # Dòng 4: GHI CHÚ ĐỎ
            worksheet.merge_cells('A4:H4')
            cell = worksheet['A4']
            cell.value = "(* Lưu ý GVBM: Chỉ nhập điểm vào các ô trống, KHÔNG chỉnh sửa cột Mã HS và Họ tên)"
            cell.font = note_font
            cell.alignment = center_align

            # --- TRANG TRÍ BẢNG DỮ LIỆU CHI TIẾT ---
            # Chỉnh độ rộng các cột cho đẹp (không bị khuất chữ)
            worksheet.column_dimensions['A'].width = 15 # Mã HS
            worksheet.column_dimensions['B'].width = 25 # Họ tên
            for col in ['C', 'D', 'E', 'F', 'G', 'H']:
                worksheet.column_dimensions[col].width = 12 # Các cột điểm

            # Tô màu và kẻ khung cho dòng Tiêu đề cột (Dòng số 6 trong Excel)
            for col_num in range(1, 9):
                cell = worksheet.cell(row=6, column=col_num)
                cell.font = bold_font
                cell.fill = header_fill
                cell.alignment = center_align
                cell.border = thin_border

            # Kẻ khung và căn lề cho tất cả các dữ liệu học sinh
            for r in range(7, 7 + len(df)):
                for c in range(1, 9):
                    cell = worksheet.cell(row=r, column=c)
                    cell.border = thin_border
                    if c > 2: # Nếu là cột điểm (Từ cột 3 trở đi) thì canh giữa
                        cell.alignment = center_align
                    else: # Nếu là Mã HS và Tên thì canh trái
                        cell.alignment = left_align

        # 3. TRẢ FILE VỀ CHO TRÌNH DUYỆT TẢI XUỐNG
        output.seek(0)
        ten_file = f"BangDiem_{mon.ten_mon_hoc}_{lop.ten_lop}_{hk.ky_hoc}.xlsx".replace(" ", "_")
        
        response = HttpResponse(
            output.read(), 
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        )
        response['Content-Disposition'] = f'attachment; filename={ten_file}'
        return response

    except Exception as e:
        messages.error(request, f"Lỗi xuất file: {e}")
        return redirect('teacher_dashboard')
    
import io
from django.http import HttpResponse
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side

@login_required(login_url='login')
def xuat_file_mau_gv(request):
    if request.user.vai_tro != 'admin':
        messages.error(request, "Bạn không có quyền tải file này!")
        return redirect('login')

    try:
        # 1. Tạo cấu trúc cột chuẩn khớp 100% với hàm import của bạn
        df = pd.DataFrame(columns=['Mã GV', 'Họ tên', 'Số điện thoại', 'Email'])
        output = io.BytesIO()

        # 2. Bắt đầu vẽ file Excel
        with pd.ExcelWriter(output, engine='openpyxl') as writer:
            df.to_excel(writer, index=False, sheet_name='Giao_Vien_Mau')
            worksheet = writer.sheets['Giao_Vien_Mau']

            # CSS cho Excel: Chữ trắng, nền xanh dương
            header_font = Font(bold=True, color='FFFFFF')
            header_fill = PatternFill(start_color="3B82F6", end_color="3B82F6", fill_type="solid")
            center_align = Alignment(horizontal='center', vertical='center')
            thin_border = Border(left=Side(style='thin'), right=Side(style='thin'), top=Side(style='thin'), bottom=Side(style='thin'))

            # Áp dụng CSS cho Tiêu đề
            for col_num, column_title in enumerate(df.columns, 1):
                cell = worksheet.cell(row=1, column=col_num)
                cell.font = header_font
                cell.fill = header_fill
                cell.alignment = center_align
                cell.border = thin_border

            # Chỉnh độ rộng các cột cho đẹp
            worksheet.column_dimensions['A'].width = 15
            worksheet.column_dimensions['B'].width = 25
            worksheet.column_dimensions['C'].width = 15
            worksheet.column_dimensions['D'].width = 30

            # 3. Chèn 1 dòng ví dụ mờ mờ cho Admin dễ hình dung
            worksheet.append(['GV001', 'Nguyễn Văn Anh', '0901234567', 'nvanh@truong.edu.vn'])
            for col in range(1, 5):
                cell = worksheet.cell(row=2, column=col)
                cell.font = Font(italic=True, color='94A3B8') # Chữ xám in nghiêng

        # 4. Trả file về cho trình duyệt
        output.seek(0)
        response = HttpResponse(
            output.read(),
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        )
        response['Content-Disposition'] = 'attachment; filename=File_Mau_GiaoVien.xlsx'
        return response

    except Exception as e:
        messages.error(request, f"Lỗi xuất file: {e}")
        return redirect('admin_dashboard')

@login_required(login_url='login')
def xuat_file_mau_hs(request):
    if request.user.vai_tro != 'admin': return redirect('login')
    try:
        df = pd.DataFrame(columns=['Mã HS', 'Họ tên', 'Email', 'Tên Lớp'])
        output = io.BytesIO()
        with pd.ExcelWriter(output, engine='openpyxl') as writer:
            df.to_excel(writer, index=False, sheet_name='Hoc_Sinh_Mau')
            worksheet = writer.sheets['Hoc_Sinh_Mau']
            header_font = Font(bold=True, color='FFFFFF')
            header_fill = PatternFill(start_color="10B981", end_color="10B981", fill_type="solid") # Màu xanh lá
            center_align = Alignment(horizontal='center', vertical='center')
            thin_border = Border(left=Side(style='thin'), right=Side(style='thin'), top=Side(style='thin'), bottom=Side(style='thin'))

            for col_num in range(1, 5):
                cell = worksheet.cell(row=1, column=col_num)
                cell.font = header_font
                cell.fill = header_fill
                cell.alignment = center_align
                cell.border = thin_border

            worksheet.column_dimensions['A'].width = 15
            worksheet.column_dimensions['B'].width = 25
            worksheet.column_dimensions['C'].width = 30
            worksheet.column_dimensions['D'].width = 15

            # Dòng ví dụ cho Admin
            worksheet.append(['HS2024001', 'Nguyễn Văn Học Sinh', 'hs001@truong.edu.vn', '10A1'])
            for col in range(1, 5):
                worksheet.cell(row=2, column=col).font = Font(italic=True, color='94A3B8')

        output.seek(0)
        response = HttpResponse(output.read(), content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
        response['Content-Disposition'] = 'attachment; filename=File_Mau_HocSinh.xlsx'
        return response
    except Exception as e:
        messages.error(request, f"Lỗi xuất file: {e}")
        return redirect('admin_dashboard')
@login_required(login_url='login')
def xuat_file_mau_ph(request):
    if request.user.vai_tro != 'admin': return redirect('login')
    try:
        df = pd.DataFrame(columns=['Số điện thoại', 'Họ tên PH', 'Mã HS liên kết'])
        output = io.BytesIO()
        with pd.ExcelWriter(output, engine='openpyxl') as writer:
            df.to_excel(writer, index=False, sheet_name='Phu_Huynh_Mau')
            worksheet = writer.sheets['Phu_Huynh_Mau']
            header_font = Font(bold=True, color='FFFFFF')
            header_fill = PatternFill(start_color="8B5CF6", end_color="8B5CF6", fill_type="solid") # Màu tím
            center_align = Alignment(horizontal='center', vertical='center')
            thin_border = Border(left=Side(style='thin'), right=Side(style='thin'), top=Side(style='thin'), bottom=Side(style='thin'))

            for col_num in range(1, 4):
                cell = worksheet.cell(row=1, column=col_num)
                cell.font = header_font
                cell.fill = header_fill
                cell.alignment = center_align
                cell.border = thin_border

            worksheet.column_dimensions['A'].width = 20
            worksheet.column_dimensions['B'].width = 25
            worksheet.column_dimensions['C'].width = 20

            # Dòng ví dụ
            worksheet.append(['0901234567', 'Phạm Văn Phụ Huynh', 'HS2024001'])
            for col in range(1, 4):
                worksheet.cell(row=2, column=col).font = Font(italic=True, color='94A3B8')

        output.seek(0)
        response = HttpResponse(output.read(), content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
        response['Content-Disposition'] = 'attachment; filename=File_Mau_PhuHuynh.xlsx'
        return response
    except Exception as e:
        messages.error(request, f"Lỗi xuất file: {e}")
        return redirect('admin_dashboard')
# ================= ĐĂNG XUẤT =================
def dang_xuat(request):
    logout(request) # Lệnh này của Django sẽ dọn sạch thẻ nhớ, đá user văng ra ngoài
    return redirect('login') # Đẩy về lại trang đăng nhập gốc